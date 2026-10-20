"""``archdoc build docx``: the Word document (SPEC 6.1).

Two modes, one block renderer:

- **Template mode** (``template:`` in arch-doc.yaml, or ``template.path`` in
  the profile): the organisation's own .docx is filled in place. Tracked
  changes and comments are cleaned; the template's Heading 1s are aligned
  with the document's chapters and its Heading 2s with each chapter's
  sections, in metamodel order; each kept section's guidance is replaced by
  content; template tables are cloned when a content table's header matches
  them; template sections and chapters the document does not carry are
  removed; every section break is kept; the document-control tables are
  filled from arch-doc.yaml; cover text named in ``template.fields`` is
  replaced; and the table of contents is rebuilt from the live headings.
- **Neutral mode** (no template): a document is generated with a cover,
  document control, a table of contents field, multilevel heading numbering,
  captions and the house table style, styled from the profile's ``house``
  palette and fonts.

Every block kind renders: paragraphs with inline bold, italic, code and
links; nested bullets and ordered lists; tables; register tables with status
chips as coloured words; figures and views (sized to the text width, capped
at the page height, captioned "Figure n. ..." and "View n. ..." with the
core's numbering); and the EA catalogue.

Figures are never faked. Each FIGURE and VIEW gets its PNG from
:func:`archdoc.diagrams.ensure_png`, which renders the Mermaid, arch-diagram
or script source, or rasterises a static SVG, when the PNG is missing or
stale; when that is impossible the build stops with :class:`DocxBuildError`
naming the figure, where it is cited and the missing tool.

:func:`read_structure` reads any .docx back (headings with their rendered
numbers, the TOC's cached entries, tables, pictures, captions, section
count); the tests and ``archdoc build`` verification use it.

Style detection (:func:`detect_styles`) and title matching
(:func:`similarity`, :func:`align`) are shared with :mod:`archdoc.extract`.
"""
from __future__ import annotations

import copy
import datetime as _dt
import difflib
import os
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml.table import CT_Tbl
from docx.shared import Emu, Pt, RGBColor, Twips
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from . import content as C
from .formatting import EA_NAMES, MONTHS, date_format, format_date, view_caption, view_note  # noqa: F401
from .project import person

EMU_PER_TWIP = 635
EMU_PER_INCH = 914400
CAPTION_ALLOWANCE = int(0.75 * EMU_PER_INCH)  # caption and spacing kept on the page with a portrait figure
CHIP_WORDS = {"done": "Done", "partial": "Partial", "not-started": "Not started", "tbc": "TBC"}
CHIP_PALETTE = {"done": "done", "partial": "partial", "not-started": "not_started", "tbc": "tbc"}
NEUTRAL_STYLES = {"h1": "Heading 1", "h2": "Heading 2", "body": "Normal", "bullet": "List Paragraph",
                  "table": "House Table", "caption": "Caption", "toc1": "toc 1", "toc2": "toc 2"}


class DocxBuildError(RuntimeError):
    """The docx cannot be built: a missing figure, a tool that is not
    installed, an unusable template. The message says what to do."""


# =================================================================== XML helpers

def tag(el) -> str:
    return el.tag.split("}")[1] if "}" in el.tag else el.tag


_T, _NBH = qn("w:t"), qn("w:noBreakHyphen")


def _chars(el):
    for x in el.iter(_T, _NBH):
        yield x, ("-" if x.tag == _NBH else (x.text or ""))


def ptext(el) -> str:
    """All text in an element, including nested runs (a non-breaking
    hyphen reads as a hyphen)."""
    return "".join(c for _, c in _chars(el))


def own_text(p) -> str:
    """A paragraph's text, excluding text boxes nested inside it."""
    out = []
    for t, ch in _chars(p):
        anc = t.getparent()
        nested = False
        while anc is not None and anc is not p:
            if tag(anc) in ("txbxContent", "fallback", "Fallback"):
                nested = True
                break
            anc = anc.getparent()
        if not nested:
            out.append(ch)
    return "".join(out)


def has_sect(el) -> bool:
    return el.find(".//" + qn("w:sectPr")) is not None


def pstyle_id(el) -> str | None:
    if tag(el) != "p":
        return None
    ps = el.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return ps.get(qn("w:val")) if ps is not None else None


def set_attr(el, name: str, value: str) -> None:
    el.set(qn(name), value)


def make(name: str, **attrs) -> Any:
    el = OxmlElement(name)
    for k, v in attrs.items():
        el.set(qn(k.replace("_", ":", 1)), str(v))
    return el


# ============================================================ style detection

@dataclass
class StyleInfo:
    """A paragraph style's effective outline level and numbering, following
    its basedOn chain."""

    style_id: str
    name: str
    outline: int | None
    num: tuple[str, str] | None  # (numId, ilvl)


def style_table(doc) -> dict[str, StyleInfo]:
    """Every paragraph style by style id, with outline level and numbering
    resolved through ``basedOn``."""
    raw = {}
    for st in doc.styles.element.findall(qn("w:style")):
        if st.get(qn("w:type")) != "paragraph":
            continue
        sid = st.get(qn("w:styleId"))
        name_el = st.find(qn("w:name"))
        based = st.find(qn("w:basedOn"))
        ppr = st.find(qn("w:pPr"))
        ol = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
        np_ = ppr.find(qn("w:numPr")) if ppr is not None else None
        num = None
        if np_ is not None:
            nid = np_.find(qn("w:numId"))
            il = np_.find(qn("w:ilvl"))
            num = (nid.get(qn("w:val")) if nid is not None else None,
                   il.get(qn("w:val")) if il is not None else None)
        raw[sid] = {"name": name_el.get(qn("w:val")) if name_el is not None else sid,
                    "based": based.get(qn("w:val")) if based is not None else None,
                    "outline": int(ol.get(qn("w:val"))) if ol is not None else None, "num": num}
    out = {}
    for sid, r in raw.items():
        outline, num, cur, seen = r["outline"], r["num"], r["based"], {sid}
        num_id, ilvl = (num or (None, None))
        while cur and cur in raw and cur not in seen and (outline is None or num_id is None or ilvl is None):
            seen.add(cur)
            b = raw[cur]
            if outline is None:
                outline = b["outline"]
            if b["num"]:
                if num_id is None:
                    num_id = b["num"][0]
                if ilvl is None:
                    ilvl = b["num"][1]
            cur = b["based"]
        out[sid] = StyleInfo(sid, r["name"], outline if outline is not None and outline < 9 else None,
                             (num_id, ilvl or "0") if num_id else None)
    return out


def paragraph_outline(p, styles: dict[str, StyleInfo]) -> int | None:
    ppr = p.find(qn("w:pPr"))
    if ppr is not None:
        ol = ppr.find(qn("w:outlineLvl"))
        if ol is not None:
            v = int(ol.get(qn("w:val")))
            return v if v < 9 else None
    sid = pstyle_id(p)
    if sid and sid in styles:
        return styles[sid].outline
    return None


def _style_name(doc, style_id: str) -> str:
    for st in doc.styles.element.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == style_id:
            n = st.find(qn("w:name"))
            return n.get(qn("w:val")) if n is not None else style_id
    return style_id


def body_paragraphs(doc) -> list:
    """Top-level body paragraphs, in order (block-level content controls
    other than the table of contents are looked into)."""
    out = []
    for el in doc.element.body.iterchildren():
        if tag(el) == "p":
            out.append(el)
        elif tag(el) == "sdt" and not is_toc_sdt(el):
            c = el.find(qn("w:sdtContent"))
            if c is not None:
                out.extend(x for x in c.iterchildren() if tag(x) == "p")
    return out


def is_toc_sdt(el) -> bool:
    return any((g.get(qn("w:val")) or "").lower().startswith("table of contents")
               for g in el.iter(qn("w:docPartGallery")))


def detect_styles(doc) -> dict[str, str]:
    """The template's style names for each role (``h1``, ``h2``, ``body``,
    ``bullet``, ``table``, ``caption``, ``toc1``, ``toc2``), detected from
    use, not from names: headings by outline level (through ``basedOn``),
    the body and bullet styles by frequency inside the chapters, the table
    style by frequency. Roles that cannot be detected are left out."""
    styles = style_table(doc)
    paras = body_paragraphs(doc)
    by_level: dict[int, Counter] = {0: Counter(), 1: Counter()}
    for p in paras:
        lvl = paragraph_outline(p, styles)
        if lvl in by_level and own_text(p).strip():
            by_level[lvl][pstyle_id(p) or "Normal"] += 1
    out: dict[str, str] = {}
    ids = {}
    for key, lvl in (("h1", 0), ("h2", 1)):
        if by_level[lvl]:
            sid = by_level[lvl].most_common(1)[0][0]
            ids[key] = sid
            out[key] = styles[sid].name if sid in styles else _style_name(doc, sid)
    h1 = ids.get("h1")
    started = False
    body, bullet = Counter(), Counter()
    for p in paras:
        sid = pstyle_id(p) or _default_pstyle(doc)
        if h1 and sid == h1 and own_text(p).strip():
            started = True
            continue
        if not started or not own_text(p).strip() or paragraph_outline(p, styles) is not None:
            continue
        numbered = p.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None or (
            sid in styles and styles[sid].num is not None)
        (bullet if numbered else body)[sid] += 1
    if body:
        out["body"] = _style_name(doc, body.most_common(1)[0][0])
    if bullet:
        out["bullet"] = _style_name(doc, bullet.most_common(1)[0][0])
    tables = Counter()
    for t in doc.element.body.iter(qn("w:tbl")):
        ts = t.find(qn("w:tblPr") + "/" + qn("w:tblStyle"))
        if ts is not None:
            tables[ts.get(qn("w:val"))] += 1
    if tables:
        out["table"] = _style_name(doc, tables.most_common(1)[0][0])
    names = {s.name.lower(): s.name for s in styles.values()}
    for key, want in (("caption", "caption"), ("toc1", "toc 1"), ("toc2", "toc 2")):
        if want in names:
            out[key] = names[want]
    return out


def _default_pstyle(doc) -> str:
    for st in doc.styles.element.findall(qn("w:style")):
        if st.get(qn("w:type")) == "paragraph" and st.get(qn("w:default")) in ("1", "true"):
            return st.get(qn("w:styleId"))
    return "Normal"


# ============================================================ title matching

_STOP = {"and", "the", "of", "for", "to", "a", "an", "in", "on", "with", "or", "by"}


def norm(s: str) -> str:
    s = s.lower().replace("&", " and ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(s.split())


def _tokens(s: str) -> set[str]:
    out = set()
    for w in norm(s).split():
        if w in _STOP:
            continue
        if len(w) > 3 and w.endswith("s"):
            w = w[:-1]
        out.add(w)
    return out


_APPENDIX_PREFIX = re.compile(r"^\s*(appendix|annex|annexure|schedule|attachment)\s+[A-Za-z0-9]{1,3}\s*[-\u2013\u2014:.)]\s*",
                              re.I)


def similarity(a: str, b: str) -> float:
    """0..1: how alike two titles are (token overlap and character
    similarity; 1.0 when they are equal after normalisation). A leading
    "Appendix A -" is ignored."""
    a, b = _APPENDIX_PREFIX.sub("", a), _APPENDIX_PREFIX.sub("", b)
    na, nb = norm(a), norm(b)
    if not na or not nb:
        return 0.0
    if na == nb:
        return 1.0
    ta, tb = _tokens(a), _tokens(b)
    jac = len(ta & tb) / len(ta | tb) if ta | tb else 0.0
    seq = difflib.SequenceMatcher(None, na, nb).ratio()
    return 0.5 * jac + 0.5 * seq


def chapter_score(title_a: str, secs_a: list[str], title_b: str, secs_b: list[str]) -> float:
    """Title similarity plus section overlap."""
    t = similarity(title_a, title_b)
    if not secs_a or not secs_b:
        return t
    hits = sum(1 for s in secs_a if max(similarity(s, x) for x in secs_b) >= 0.6)
    overlap = hits / len(secs_a)
    return 0.55 * t + 0.45 * overlap


def align(n: int, m: int, score: Callable[[int, int], float], threshold: float) -> list[tuple[int, int]]:
    """An order-preserving one-to-one matching of ``range(n)`` to
    ``range(m)`` that maximises the total score of pairs scoring at least
    ``threshold``."""
    s = [[score(i, j) for j in range(m)] for i in range(n)]
    dp = [[0.0] * (m + 1) for _ in range(n + 1)]
    for i in range(n - 1, -1, -1):
        for j in range(m - 1, -1, -1):
            best = max(dp[i + 1][j], dp[i][j + 1])
            if s[i][j] >= threshold:
                best = max(best, s[i][j] + dp[i + 1][j + 1])
            dp[i][j] = best
    pairs, i, j = [], 0, 0
    while i < n and j < m:
        if s[i][j] >= threshold and abs(dp[i][j] - (s[i][j] + dp[i + 1][j + 1])) < 1e-9:
            pairs.append((i, j))
            i, j = i + 1, j + 1
        elif dp[i][j] == dp[i + 1][j]:
            i += 1
        else:
            j += 1
    return pairs


def fill_gaps(pairs: list[tuple[int, int]], n: int, m: int, score: Callable[[int, int], float],
              floor: float) -> list[tuple[int, int]]:
    """Add pairs by position: where the items between two matched
    neighbours are as many on both sides, pair them in order when each
    scores at least ``floor``. A renamed chapter between two that matched
    is the same chapter."""
    out = sorted(pairs)
    bounds = [(-1, -1)] + out + [(n, m)]
    added = []
    for (i1, j1), (i2, j2) in zip(bounds, bounds[1:]):
        k = i2 - i1 - 1
        if k <= 0 or k != j2 - j1 - 1:
            continue
        cand = [(i1 + d, j1 + d) for d in range(1, k + 1)]
        if all(score(i, j) >= floor for i, j in cand):
            added += cand
    return sorted(out + added)


def match(n: int, m: int, score: Callable[[int, int], float], threshold: float,
          floor: float = 0.1) -> list[tuple[int, int]]:
    """:func:`align` by score, then :func:`fill_gaps` by position."""
    return fill_gaps(align(n, m, score, threshold), n, m, score, floor)


# ============================================================ numbering model

class Numbering:
    """Reads and extends the numbering part: abstract definitions, num
    instances, and new bullet / ordered / heading definitions."""

    def __init__(self, doc):
        self.doc = doc
        self.el = doc.part.numbering_part.element

    def _next(self, name: str, attr: str) -> int:
        vals = [int(x.get(qn(attr))) for x in self.el.findall(qn(name)) if (x.get(qn(attr)) or "").isdigit()]
        return max(vals, default=0) + 1

    def add_abstract(self, xml_levels: str, multi: str = "hybridMultilevel") -> str:
        aid = str(self._next("w:abstractNum", "w:abstractNumId"))
        an = parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid}">'
                       f'<w:multiLevelType w:val="{multi}"/>{xml_levels}</w:abstractNum>')
        first_num = self.el.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(an)
        else:
            self.el.append(an)
        return aid

    def add_num(self, abstract_id: str, restart: bool = False) -> str:
        nid = str(self._next("w:num", "w:numId"))
        override = '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>' if restart else ""
        num = parse_xml(f'<w:num {nsdecls("w")} w:numId="{nid}"><w:abstractNumId w:val="{abstract_id}"/>'
                        f'{override}</w:num>')
        cleanup = self.el.find(qn("w:numIdMacAtCleanup"))
        if cleanup is not None:
            cleanup.addprevious(num)
        else:
            self.el.append(num)
        return nid

    def abstract_of(self, num_id: str) -> Any:
        for n in self.el.findall(qn("w:num")):
            if n.get(qn("w:numId")) == str(num_id):
                aid = n.find(qn("w:abstractNumId")).get(qn("w:val"))
                for a in self.el.findall(qn("w:abstractNum")):
                    if a.get(qn("w:abstractNumId")) == aid:
                        return a
        return None

    def level_format(self, num_id: str, ilvl: int = 0) -> str | None:
        a = self.abstract_of(num_id)
        if a is None:
            return None
        for lvl in a.findall(qn("w:lvl")):
            if lvl.get(qn("w:ilvl")) == str(ilvl):
                f = lvl.find(qn("w:numFmt"))
                return f.get(qn("w:val")) if f is not None else None
        return None

    def bullet_definition(self, font: str) -> str:
        chars = ["•", "◦", "▪"]
        levels = "".join(
            f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
            f'<w:lvlText w:val="{chars[i % 3]}"/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:ind w:left="{360 * (i + 1)}" w:hanging="360"/></w:pPr>'
            f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}"/></w:rPr></w:lvl>' for i in range(9))
        return self.add_num(self.add_abstract(levels))

    def ordered_abstract(self) -> str:
        fmts = ["decimal", "lowerLetter", "lowerRoman"]
        levels = "".join(
            f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmts[i % 3]}"/>'
            f'<w:lvlText w:val="%{i + 1}."/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:ind w:left="{360 * (i + 1)}" w:hanging="360"/></w:pPr></w:lvl>' for i in range(9))
        return self.add_abstract(levels)

    def heading_definition(self, style_ids: list[str]) -> str:
        levels = []
        for i in range(9):
            text = ".".join(f"%{k + 1}" for k in range(i + 1))
            ps = f'<w:pStyle w:val="{style_ids[i]}"/>' if i < len(style_ids) else ""
            ind = 567 + 142 * i
            levels.append(f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="decimal"/>{ps}'
                          f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
                          f'<w:pPr><w:ind w:left="{ind}" w:hanging="{ind}"/></w:pPr></w:lvl>')
        return self.add_num(self.add_abstract("".join(levels), "multilevel"))


# ============================================================ catalogue

SHARED_KEYS = ("shared_with", "inheritedFrom", "inherited_from")


def catalogue_table(project, cat: C.Catalogue) -> tuple[list[str], list[list[str]]]:
    """The EA catalogue as a table: the core's object counts by type, with
    two optional columns when the model carries the data: objects shared
    from another design (``shared_with`` / ``inheritedFrom`` on an object),
    and examples (the first object names). arch-doc.yaml
    ``catalogue.columns`` limits them (``[shared, examples]``) and
    ``catalogue.labels`` names the headers (``type``, ``objects``,
    ``shared``, ``examples``; placeholders resolve)."""
    header, rows = list(cat.header), [list(r) for r in cat.body]
    path = project.root / cat.model_path
    if not cat.found or not path.exists():
        return header, rows
    import json
    objects = json.loads(path.read_text()).get("objects", [])
    want = (project.config.get("catalogue") or {}).get("columns") if isinstance(
        project.config.get("catalogue"), dict) else None
    want = set(want) if want is not None else {"shared", "examples"}
    by_type: dict[str, list[dict]] = {}
    for o in objects:
        by_type.setdefault(o.get("type", "unknown"), []).append(o)
    shared_of = lambda o: next((o[k] for k in SHARED_KEYS if o.get(k)), None)  # noqa: E731
    shared = "shared" in want and any(shared_of(o) for o in objects)
    examples = "examples" in want and any(o.get("name") for o in objects)
    labels = (project.config.get("catalogue") or {}).get("labels") if isinstance(
        project.config.get("catalogue"), dict) else None
    labels = {k: project.resolve_text(str(v))[0] for k, v in (labels or {}).items()} if isinstance(labels, dict) else {}
    if not (shared or examples):
        return [labels.get("type", header[0]), labels.get("objects", header[1])], rows
    sources = {str(shared_of(o)) for o in objects if shared_of(o)}
    if shared:
        header = [labels.get("type", header[0]), labels.get("objects", "Objects in this design"),
                  labels.get("shared", f"Shared with the {sources.pop()}" if len(sources) == 1
                             else "Shared from other designs")]
    else:
        header = [labels.get("type", header[0]), labels.get("objects", header[1])]
    if examples:
        header.append(labels.get("examples", "Examples"))
    out = []
    for row in rows:
        objs = by_type.get(row[0], [])
        new = [row[0]]
        own = [o for o in objs if not shared_of(o)]
        if shared:
            new += [str(len(own)), str(len(objs) - len(own))]
        else:
            new.append(row[1])
        if examples:
            new.append("; ".join(o["name"] for o in (own or objs)[:2] if o.get("name")))
        out.append(new)
    return header, out


# ============================================================ view captions

# view_caption, view_note and EA_NAMES live in archdoc.formatting (imported above).


# ============================================================ figure files

INSTALL_HINT = {
    "mermaid": "install Node.js so npx can run @mermaid-js/mermaid-cli",
    "arch-diagram": "install the arch-diagram plugin, or render the PNG with it elsewhere",
    "chrome": "install Chrome or Chromium to rasterise the SVG",
}


def figure_png(fig: C.Figure, project) -> Path:
    """The PNG to embed for a FIGURE or VIEW, from
    :func:`archdoc.diagrams.ensure_png` (which renders or rasterises it when
    needed). Raises DocxBuildError naming the figure, where it is cited and
    what is missing; there is never a placeholder image."""
    from . import diagrams
    where = f"{fig.source}: " if fig.source else ""
    if fig.path is None and fig.source_file is None:
        raise DocxBuildError(f"{where}{fig.label} '{fig.name}' has no file: add diagrams/{fig.name}.png, "
                             f".svg or a source (.mmd, .ead, .py)")
    try:
        png = diagrams.ensure_png(fig, project)
    except diagrams.DiagramError as e:
        hint = f"; {INSTALL_HINT[e.tool]}" if getattr(e, "tool", None) in INSTALL_HINT else ""
        raise DocxBuildError(f"{where}{fig.label} '{fig.name}' cannot be embedded: {e}{hint}") from None
    if not png.exists():
        raise DocxBuildError(f"{where}{fig.label} file {png} is missing")
    return png


def image_size(path: Path) -> tuple[int, int]:
    from PIL import Image
    with Image.open(path) as im:
        return im.size


# ============================================================ page geometry

def governing_sectpr(el, body) -> Any:
    """The sectPr that governs where ``el`` sits: the next paragraph-level
    sectPr after it, or the body's final one."""
    cur = el
    while cur is not None and cur.getparent() is not body:
        cur = cur.getparent()
    nxt = cur
    while nxt is not None:
        if tag(nxt) == "sectPr":
            return nxt
        if tag(nxt) == "p":
            s = nxt.find(qn("w:pPr") + "/" + qn("w:sectPr"))
            if s is not None:
                return s
        nxt = nxt.getnext()
    return body.find(qn("w:sectPr"))


def text_box(sect) -> tuple[int, int]:
    """(text width, text height) in EMU for a sectPr."""
    pg = sect.find(qn("w:pgSz")) if sect is not None else None
    mar = sect.find(qn("w:pgMar")) if sect is not None else None
    w = int(pg.get(qn("w:w"))) if pg is not None and pg.get(qn("w:w")) else 11906
    h = int(pg.get(qn("w:h"))) if pg is not None and pg.get(qn("w:h")) else 16838

    def m(name, default):
        if mar is None or mar.get(qn(name)) is None:
            return default
        return abs(int(float(mar.get(qn(name)))))
    tw = w - m("w:left", 1440) - m("w:right", 1440)
    th = h - m("w:top", 1440) - m("w:bottom", 1440)
    return tw * EMU_PER_TWIP, th * EMU_PER_TWIP


def fit_figure(px_w: int, px_h: int, box: tuple[int, int]) -> int:
    """The EMU width for a figure: the text width, reduced so a tall figure
    and its caption fit on the page."""
    tw, th = box
    width = tw
    max_h = th - CAPTION_ALLOWANCE
    if px_w and px_h * width / px_w > max_h:
        width = int(max_h * px_w / px_h)
    return int(width)


# ============================================================ the renderer

def _hex(c: str) -> str:
    return c.lstrip("#").upper()


@dataclass
class Styles:
    """Resolved style ids for each role."""

    h1: str
    h2: str
    body: str
    bullet: str
    table: str | None
    caption: str
    toc1: str
    toc2: str


class Renderer:
    """Turns content blocks into body elements (``w:p`` and ``w:tbl``) for
    one document, in either mode."""

    def __init__(self, project, doc, styles: Styles, numbering: Numbering,
                 bullet_num: str, ordered_abstract: str, prototypes: Callable[[], list] | None = None):
        self.p = project
        self.doc = doc
        self.body = doc.element.body
        self.st = styles
        self.num = numbering
        self.bullet_num = bullet_num
        self.ordered_abstract = ordered_abstract
        self.palette = project.profile.palette
        self.fonts = project.profile.fonts
        self.figure_count = 0
        self.view_count = 0
        self.squeezed = []

    # -------------------------------------------------------------- runs
    def para(self, style_id: str | None = None) -> Paragraph:
        p_el = OxmlElement("w:p")
        if style_id:
            ppr = p_el.get_or_add_pPr()
            ppr.get_or_add_pStyle().set(qn("w:val"), style_id)
        return Paragraph(p_el, self.doc)

    def runs(self, par: Paragraph, text: str, base_rpr=None) -> None:
        """Inline Markdown into runs: bold, italic, code (the house mono
        font) and links (external hyperlinks)."""
        for chunk, styles, href in C.inline_runs(text):
            if not chunk:
                continue
            r = par.add_run(chunk)
            if base_rpr is not None:
                r._r.insert(0, copy.deepcopy(base_rpr))
            if "bold" in styles:
                r.bold = True
            if "em" in styles:
                r.italic = True
            if "code" in styles:
                mono = self.fonts.get("mono") or "Consolas"
                r.font.name = mono
                rf = r._r.get_or_add_rPr().get_or_add_rFonts()
                rf.set(qn("w:cs"), mono)
                rf.set(qn("w:eastAsia"), mono)
            no_break_identifiers(r._r)
            if href:
                self._hyperlink(par, r, href)

    def _hyperlink(self, par: Paragraph, run: Run, href: str) -> None:
        hl = OxmlElement("w:hyperlink")
        if href.startswith("#"):
            hl.set(qn("w:anchor"), href[1:])
        else:
            hl.set(qn("r:id"), self.doc.part.relate_to(href, RT.HYPERLINK, is_external=True))
        r_el = run._r
        par._p.replace(r_el, hl) if r_el.getparent() is par._p else None
        hl.append(r_el)
        if self._has_style("Hyperlink", "character"):
            run.style = self.doc.styles["Hyperlink"]
        else:
            run.font.underline = True
            run.font.color.rgb = RGBColor.from_string(_hex(self.palette["accent"]))

    def _has_style(self, name: str, kind: str) -> bool:
        for st in self.doc.styles.element.findall(qn("w:style")):
            n = st.find(qn("w:name"))
            if st.get(qn("w:type")) == kind and n is not None and n.get(qn("w:val")) == name:
                return True
        return False

    # ------------------------------------------------------------ blocks
    def blocks(self, blocks: Iterable[C.Block], anchor) -> list:
        """Elements for blocks; ``anchor`` is the element they will follow
        (for page geometry)."""
        out = []
        for b in blocks:
            out.extend(self.block(b, anchor))
        return out

    def block(self, b: C.Block, anchor) -> list:
        if isinstance(b, C.Para):
            par = self.para(self.st.body)
            self.runs(par, b.text)
            return [par._p]
        if isinstance(b, C.Bullets):
            return self.bullets(b)
        if isinstance(b, C.View):
            return self.figure(b, anchor, "View")
        if isinstance(b, C.Figure):
            return self.figure(b, anchor, "Figure")
        if isinstance(b, C.RegisterTable):
            return self.register(b, anchor)
        if isinstance(b, C.Catalogue):
            header, rows = catalogue_table(self.p, b)
            return self.table(header, rows, anchor, where=describe_block(self.p, b, "the EA catalogue"))
        if isinstance(b, C.Table):
            return self.table(b.header, b.body, anchor, where=describe_block(self.p, b, "the table"))
        raise DocxBuildError(f"{b.source}: no docx rendering for block kind '{b.kind}'")

    def bullets(self, b: C.Bullets) -> list:
        num_id = self.num.add_num(self.ordered_abstract, restart=True) if b.ordered else self.bullet_num
        out = []
        for text, level in zip(b.items, b.levels):
            par = self.para(self.st.bullet)
            numpr = par._p.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_ilvl().val = min(level, 8)
            numpr.get_or_add_numId().val = int(num_id)
            self.runs(par, text)
            out.append(par._p)
        return out

    def caption(self, label: str, number: int, text: str) -> Any:
        """'Figure n. text' with a SEQ field whose cached result is the
        core's number, so Word renumbers consistently on update."""
        par = self.para(self.st.caption)
        par.add_run(f"{label} ")
        self._field(par, f" SEQ {label} \\* ARABIC ", str(number))
        self.runs(par, ". " + text)
        return par._p

    def _field(self, par: Paragraph, instr: str, result: str) -> None:
        r1 = par.add_run()
        r1._r.append(make("w:fldChar", w_fldCharType="begin"))
        r2 = par.add_run()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        r2._r.append(it)
        r3 = par.add_run()
        r3._r.append(make("w:fldChar", w_fldCharType="separate"))
        par.add_run(result)
        r5 = par.add_run()
        r5._r.append(make("w:fldChar", w_fldCharType="end"))

    def figure(self, f: C.Figure, anchor, label: str) -> list:
        png = figure_png(f, self.p)
        pw, ph = image_size(png)
        box = text_box(governing_sectpr(anchor, self.body)) if anchor is not None else \
            text_box(self.body.find(qn("w:sectPr")))
        width = fit_figure(pw, ph, box)
        pic = self.para(self.st.body)
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.keep_with_next = True
        pic.add_run().add_picture(str(png), width=Emu(width))
        caption = view_caption(self.p, f) if isinstance(f, C.View) else f.caption
        return [pic._p, self.caption(label, f.number, caption)]

    # ------------------------------------------------------------ tables
    added_sections = 0
    squeezed: list[str] = []
    allow_landscape = True  # neutral mode only; a template's section layout is never changed

    def table(self, header: list[str], rows: list[list[str]], anchor, chips: list[str] | None = None,
              status_col: int | None = None, empty_note: str | None = None, where: str | None = None) -> list:
        """The table's elements: the table, fitted to its content (in its
        own landscape section when it cannot fit the page otherwise), then
        a small spacer so the next paragraph does not touch it."""
        tbl = self.prototype_table(header) if self.prototypes_available() else None
        cloned = tbl is not None
        if cloned:
            self._fill_prototype(tbl, header, rows, chips, status_col, empty_note)
        else:
            tbl = self._new_table(header, rows, anchor, chips, status_col, empty_note)
        sect = governing_sectpr(anchor, self.body) if anchor is not None else self.body.find(qn("w:sectPr"))
        out = [tbl]
        result = fit_table(self.doc, tbl, text_box(sect)[0] // EMU_PER_TWIP, keep_if_fits=cloned,
                           force=not self.allow_landscape)
        if result == "landscape":
            land = landscape_of(sect)
            result = fit_table(self.doc, tbl, text_box(land)[0] // EMU_PER_TWIP, force=True)
            out = [section_break(sect), tbl, section_break(land)]
            self.added_sections += 2
        if result == "squeezed":
            self.squeezed.append(squeeze_warning(where or "a table", header, self.allow_landscape))
        finish_header(self.doc, tbl)
        out.append(self.spacer())
        return out

    def spacer(self):
        """A short empty paragraph after a table."""
        p = self.para(self.st.body)
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(0)
        pf.line_spacing = Pt(6)
        pf.line_spacing_rule = 4  # exactly
        return p._p

    def prototypes_available(self) -> bool:
        return False

    def prototype_table(self, header: list[str]):
        return None

    def _new_table(self, header, rows, anchor, chips, status_col, empty_note):
        ncols = max(1, len(header))
        box = text_box(governing_sectpr(anchor, self.body)) if anchor is not None else \
            text_box(self.body.find(qn("w:sectPr")))
        width = box[0]
        tbl_el = CT_Tbl.new_tbl(1 + max(len(rows), 1 if empty_note else 0), ncols, Emu(width))
        t = Table(tbl_el, self.doc)
        if self.st.table:
            t.style = self.doc.styles[_style_name(self.doc, self.st.table)]
        tblpr = tbl_el.tblPr
        look = tblpr.find(qn("w:tblLook"))
        if look is None:
            look = OxmlElement("w:tblLook")
            tblpr.append(look)
        for k, v in (("w:val", "04A0"), ("w:firstRow", "1"), ("w:lastRow", "0"), ("w:firstColumn", "0"),
                     ("w:lastColumn", "0"), ("w:noHBand", "0"), ("w:noVBand", "1")):
            look.set(qn(k), v)
        all_rows = [header] + rows
        for ri, row in enumerate(t.rows):
            cells = row.cells
            for ci, cell in enumerate(cells):
                if ri < len(all_rows):
                    val = all_rows[ri][ci] if ci < len(all_rows[ri]) else ""
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_before = Pt(1)
                    par.paragraph_format.space_after = Pt(1)
                    self.runs(par, val)
                    if ri > 0 and status_col == ci and chips:
                        self._chip(par, chips[ri - 1])
        if not rows and empty_note:
            r = t.rows[1]
            merged = r.cells[0].merge(r.cells[-1]) if ncols > 1 else r.cells[0]
            par = merged.paragraphs[0]
            for extra in merged.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
            for run in list(par.runs):
                run._r.getparent().remove(run._r)
            run = par.add_run(empty_note)
            run.italic = True
        return tbl_el

    def _chip(self, par: Paragraph, chip: str) -> None:
        colour = self.palette.get(CHIP_PALETTE.get(chip, "tbc"))
        for run in par.runs:
            run.bold = True
            if colour:
                run.font.color.rgb = RGBColor.from_string(_hex(colour))

    def _fill_prototype(self, tbl, header, rows, chips, status_col, empty_note) -> None:
        trs = tbl.findall(qn("w:tr"))
        hdr, body_proto = trs[0], copy.deepcopy(trs[1] if len(trs) > 1 else trs[0])
        for tr in trs[1:]:
            tbl.remove(tr)
        self._set_row(hdr, header)
        for ri, row in enumerate(rows):
            tr = copy.deepcopy(body_proto)
            self._set_row(tr, row)
            if status_col is not None and chips:
                tcs = tr.findall(qn("w:tc"))
                if status_col < len(tcs):
                    for p_el in tcs[status_col].findall(qn("w:p")):
                        self._chip(Paragraph(p_el, self.doc), chips[ri])
            tbl.append(tr)
        if not rows and empty_note:
            tr = copy.deepcopy(body_proto)
            self._set_row(tr, [empty_note] + [""] * (len(tr.findall(qn("w:tc"))) - 1))
            tbl.append(tr)

    def _set_row(self, tr, values: list[str]) -> None:
        tcs = tr.findall(qn("w:tc"))
        for i, tc in enumerate(tcs):
            self.set_cell(tc, values[i] if i < len(values) else "")

    def set_cell(self, tc, text: str) -> None:
        """Replace a cell's text, keeping its first paragraph's properties
        and first run's formatting."""
        ps = tc.findall(qn("w:p"))
        if not ps:
            ps = [OxmlElement("w:p")]
            tc.append(ps[0])
        p = ps[0]
        for extra in ps[1:]:
            tc.remove(extra)
        first_rpr = p.find(".//" + qn("w:r") + "/" + qn("w:rPr"))
        rpr = keep_cell_rpr(copy.deepcopy(first_rpr)) if first_rpr is not None else None
        for ch in list(p):
            if tag(ch) != "pPr":
                p.remove(ch)
        self.runs(Paragraph(p, self.doc), text, rpr)

    def register(self, b: C.RegisterTable, anchor) -> Any:
        """A register table from the core's compact rows: labelled values,
        and an inherited TBC's owner named once per row rather than in every
        cell."""
        status_col = b.columns.index("status") if "status" in b.columns else None
        rows = b.compact_rows if len(b.compact_rows) == len(b.body) else b.body
        return self.table(b.header, [list(r) for r in rows], anchor, b.chips, status_col, "None recorded.",
                          where=describe_block(self.p, b, f"the {b.type} register"))


# ============================================================ identifiers

# A token that is an identifier rather than a compound word: it has a digit
# (RSK-04, IA-2026-07, 2027-02-01) or three or more hyphen-joined parts
# (pts-prd-tracking-db). Its hyphens never break a line.
_IDENT = re.compile(r"(?<!\S)(?=[^\s]*\d)[^\s-]+(?:-[^\s-]+)+|(?<!\S)[^\s-]+(?:-[^\s-]+){2,}")


def identifier_tokens(text: str) -> list[str]:
    return [m.group(0) for m in _IDENT.finditer(text)]


def no_break_identifiers(r) -> None:
    """Replace the hyphens inside identifier tokens of a run with
    ``w:noBreakHyphen``, so an id such as RSK-04 never splits at a line end."""
    for t in list(r.findall(qn("w:t"))):
        text = t.text or ""
        spans = [m.span() for m in _IDENT.finditer(text)]
        if not spans:
            continue
        parts: list[Any] = []
        pos = 0
        for a, b in spans:
            seg = text[a:b].split("-")
            parts.append(text[pos:a] + seg[0])
            for piece in seg[1:]:
                parts.append(None)
                parts.append(piece)
            pos = b
        parts[-1] = (parts[-1] or "") + text[pos:]
        anchor = t
        first = True
        for piece in parts:
            if piece is None:
                el = OxmlElement("w:noBreakHyphen")
            else:
                if first:
                    t.text = piece
                    first = False
                    continue
                el = OxmlElement("w:t")
                el.set(qn("xml:space"), "preserve")
                el.text = piece
            anchor.addnext(el)
            anchor = el


# ============================================================ table headers and warnings

def describe_block(project, b, what: str) -> str:
    """Where a block is, for a message: its anchor, section and what it is."""
    anchor = b.anchor or ""
    m = re.match(r"^s(\d+)(?:-(\d+))?", anchor)
    place = ""
    if m:
        node = project.by_number(f"{m.group(1)}.{m.group(2)}" if m.group(2) else m.group(1))
        if node is not None:
            place = f" in {getattr(node, 'number', None) or node.number} {node.title}"
    src = f" ({b.source})" if b.source else ""
    return f"{what} {anchor}{place}{src}".strip()


def squeeze_warning(where: str, header: list[str], neutral: bool) -> str:
    cols = len(header)
    fix = ("give the REGISTER fewer columns (cols ...)" if " register " in f" {where} "
           else "give the table fewer columns or split it")
    if neutral:
        return (f"{where}: its {cols} columns do not fit even at {MIN_TABLE_PT:g}pt on a landscape page, so a "
                f"long word may break; {fix}")
    return (f"{where}: its {cols} columns do not fit the template's page width at {MIN_TABLE_PT:g}pt, and a "
            f"template's page layout is kept, so a long word may break; {fix}")


_TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark"]


def _first_row_format(doc, tbl):
    """(rPr children, cell shd) the table style gives its first row."""
    ts = tbl.find(qn("w:tblPr") + "/" + qn("w:tblStyle"))
    sid = ts.get(qn("w:val")) if ts is not None else None
    rpr_kids, shd, seen = {}, None, set()
    while sid and sid not in seen:
        seen.add(sid)
        st = _style_el(doc, sid)
        if st is None:
            break
        for sp in st.findall(qn("w:tblStylePr")):
            if sp.get(qn("w:type")) != "firstRow":
                continue
            rpr = sp.find(qn("w:rPr"))
            for ch in (rpr if rpr is not None else []):
                rpr_kids.setdefault(tag(ch), ch)
            s_ = sp.find(qn("w:tcPr") + "/" + qn("w:shd"))
            if s_ is not None and shd is None:
                shd = s_
        b = st.find(qn("w:basedOn"))
        sid = b.get(qn("w:val")) if b is not None else None
    return rpr_kids, shd


def finish_header(doc, tbl) -> None:
    """Make the first row a header that repeats on every page and keeps
    with the first data row. The table style's first-row formatting is
    written into the row as well, because LibreOffice drops a style's
    first-row formatting from a row that keeps with the next."""
    trs = tbl.findall(qn("w:tr"))
    if len(trs) < 2:
        return
    tr = trs[0]
    trpr = tr.find(qn("w:trPr"))
    if trpr is None:
        trpr = OxmlElement("w:trPr")
        first_tc = tr.find(qn("w:tc"))
        first_tc.addprevious(trpr) if first_tc is not None else tr.append(trpr)
    if trpr.find(qn("w:tblHeader")) is None:
        trpr.append(OxmlElement("w:tblHeader"))
    look = tbl.find(qn("w:tblPr") + "/" + qn("w:tblLook"))
    first_on = True  # the table style's first-row formatting applies unless tblLook switches it off
    if look is not None:
        if look.get(qn("w:firstRow")) is not None:
            first_on = look.get(qn("w:firstRow")) in ("1", "true")
        elif (look.get(qn("w:val")) or "").strip():
            first_on = bool(int(look.get(qn("w:val")), 16) & 0x20)
    rpr_kids, shd = _first_row_format(doc, tbl) if first_on else ({}, None)
    for tc in tr.findall(qn("w:tc")):
        if shd is not None:
            tcpr = tc.find(qn("w:tcPr"))
            if tcpr is None:
                tcpr = OxmlElement("w:tcPr")
                tc.insert(0, tcpr)
            if tcpr.find(qn("w:shd")) is None:
                tcpr.append(copy.deepcopy(shd))
                kids = sorted(tcpr, key=lambda e: _TCPR_ORDER.index(tag(e)) if tag(e) in _TCPR_ORDER else 99)
                for k in kids:
                    tcpr.append(k)
        for p in tc.findall(qn("w:p")):
            Paragraph(p, None).paragraph_format.keep_with_next = True
            for r in p.iter(qn("w:r")):
                if not rpr_kids:
                    break
                rpr = r.find(qn("w:rPr"))
                if rpr is None:
                    rpr = OxmlElement("w:rPr")
                    r.insert(0, rpr)
                have = {tag(c) for c in rpr}
                for name, el in rpr_kids.items():
                    if name not in have:
                        rpr.append(copy.deepcopy(el))
                _order_rpr(rpr)


# ============================================================ table fitting

# Advance widths in thousandths of an em for a Helvetica-like face; text in
# a table is measured with these, times FIT_SAFETY for wider substitutes.
_CHAR_W = {}
for _group, _w in (("ijl", 222), (" .,:;!I'|/ft", 278), ("r-()[]\"*", 333), ("Jcksvxyz", 500),
                   ("abdeghnopqu0123456789_?#$", 556), ("+=<>~", 584), ("FTZL", 611),
                   ("ABEKPSVXY&", 667), ("CDHNRUw", 722), ("GOQ", 778), ("Mm%", 833), ("W", 944), ("@", 1015)):
    for _c in _group:
        _CHAR_W[_c] = _w
FIT_SAFETY = 1.3
BOLD = 1.12
MIN_TABLE_PT = 8.0
DEFAULT_CELL_PAD = 216  # twips: Word's 0.075in each side, rounded up


def text_width(text: str, pt: float, bold: bool = False) -> float:
    """Width of a line of text in twips at ``pt`` points."""
    em = sum(_CHAR_W.get(c, 600) for c in text) / 1000
    return em * pt * 20 * FIT_SAFETY * (BOLD if bold else 1.0)


def _sz(el) -> float | None:
    if el is None:
        return None
    s = el.find(qn("w:sz"))
    return int(s.get(qn("w:val"))) / 2 if s is not None and s.get(qn("w:val")) else None


def _style_el(doc, style_id: str | None):
    if not style_id:
        return None
    for st in doc.styles.element.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == style_id:
            return st
    return None


def table_font_pt(doc, tbl) -> float:
    """The point size table text renders at: the table style's, else the
    default paragraph style's, else the document default."""
    ts = tbl.find(qn("w:tblPr") + "/" + qn("w:tblStyle"))
    sid = ts.get(qn("w:val")) if ts is not None else None
    seen = set()
    while sid and sid not in seen:
        seen.add(sid)
        st = _style_el(doc, sid)
        if st is None:
            break
        v = _sz(st.find(qn("w:rPr")))
        if v:
            return v
        b = st.find(qn("w:basedOn"))
        sid = b.get(qn("w:val")) if b is not None else None
    v = _sz(_style_el(doc, _default_pstyle(doc)).find(qn("w:rPr"))) if _style_el(doc, _default_pstyle(doc)) is not None else None
    if v:
        return v
    d = doc.styles.element.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
    return _sz(d) or 10.0


def cell_padding(doc, tbl) -> int:
    """Left plus right cell margin in twips."""
    def from_mar(mar):
        if mar is None:
            return None
        tot, found = 0, False
        for side in ("left", "start", "right", "end"):
            el = mar.find(qn(f"w:{side}"))
            if el is not None and el.get(qn("w:w")):
                tot += int(el.get(qn("w:w")))
                found = True
        return tot if found else None
    v = from_mar(tbl.find(qn("w:tblPr") + "/" + qn("w:tblCellMar")))
    if v is None:
        ts = tbl.find(qn("w:tblPr") + "/" + qn("w:tblStyle"))
        st = _style_el(doc, ts.get(qn("w:val"))) if ts is not None else None
        if st is not None:
            v = from_mar(st.find(qn("w:tblPr") + "/" + qn("w:tblCellMar")))
    return (v if v is not None else DEFAULT_CELL_PAD) + 60


def _cells(tbl):
    """Rows of (column index, span, text, bold, explicit pt) per cell."""
    out = []
    for ri, tr in enumerate(tbl.findall(qn("w:tr"))):
        row, col = [], 0
        for tc in tr.findall(qn("w:tc")):
            span_el = tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan"))
            span = int(span_el.get(qn("w:val"))) if span_el is not None else 1
            runs = list(tc.iter(qn("w:r")))
            bold = ri == 0 or any(r.find(qn("w:rPr") + "/" + qn("w:b")) is not None for r in runs)
            pts = [_sz(r.find(qn("w:rPr"))) for r in runs]
            pt = max((p for p in pts if p), default=None)
            texts = [ptext(p) for p in tc.findall(qn("w:p"))]
            row.append((col, span, texts, bold, pt))
            col += span
        out.append(row)
    return out


def column_needs(doc, tbl, pt: float) -> tuple[list[float], list[float]]:
    """(minimum, natural) width per grid column in twips at ``pt``: the
    minimum is the widest unbreakable token, the natural width the longest
    line unwrapped."""
    pad = cell_padding(doc, tbl)
    base = table_font_pt(doc, tbl)
    rows = _cells(tbl)
    n = max((c + s for r in rows for c, s, *_ in r), default=0)
    grid = tbl.find(qn("w:tblGrid"))
    n = max(n, len(grid.findall(qn("w:gridCol"))) if grid is not None else 0)
    mins, nats = [float(pad)] * n, [float(pad)] * n
    for row in rows:
        for col, span, texts, bold, cpt in row:
            if span != 1:
                continue
            size = pt * (cpt / base) if cpt else pt
            for line in texts:
                for tok in line.split():
                    mins[col] = max(mins[col], text_width(tok, size, bold) + pad)
                nats[col] = max(nats[col], text_width(line, size, bold) + pad)
    return mins, nats


def _set_widths(tbl, widths: list[int], total: int) -> None:
    tblpr = tbl.find(qn("w:tblPr"))
    tw = tblpr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tblpr.insert_element_before(tw, "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                                    "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
                                    "w:tblDescription")
    tw.set(qn("w:w"), str(total))
    tw.set(qn("w:type"), "dxa")
    lay = tblpr.find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout")
        tblpr.insert_element_before(lay, "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")
    lay.set(qn("w:type"), "fixed")
    grid = tbl.find(qn("w:tblGrid"))
    for gc in grid.findall(qn("w:gridCol")):
        grid.remove(gc)
    for w in widths:
        grid.append(make("w:gridCol", w_w=w))
    trs = tbl.findall(qn("w:tr"))
    # Rows are not made unsplittable: LibreOffice drops the body of a table whose repeating header
    # lands at a page foot. The house table keeps its header with the first row (see _new_table).
    for tr in trs:
        col = 0
        for tc in tr.findall(qn("w:tc")):
            tcpr = tc.find(qn("w:tcPr"))
            if tcpr is None:
                tcpr = OxmlElement("w:tcPr")
                tc.insert(0, tcpr)
            span_el = tcpr.find(qn("w:gridSpan"))
            span = int(span_el.get(qn("w:val"))) if span_el is not None else 1
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.insert(0, tcw)
            tcw.set(qn("w:w"), str(sum(widths[col:col + span])))
            tcw.set(qn("w:type"), "dxa")
            col += span


def _set_table_pt(tbl, pt: float, base: float) -> None:
    """Scale every run in the table to ``pt`` (runs with their own size
    keep their ratio to the table's size)."""
    for r in tbl.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            r.insert(0, rpr)
        own = _sz(rpr)
        size = pt * (own / base) if own else pt
        for name in ("w:sz", "w:szCs"):
            el = rpr.find(qn(name))
            if el is None:
                el = OxmlElement(name)
                rpr.append(el)
            el.set(qn("w:val"), str(int(round(size * 2))))
    # the schema orders rPr children; sz/szCs appended last are reordered here
    for rpr in tbl.iter(qn("w:rPr")):
        _order_rpr(rpr)


_RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
              "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
              "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
              "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]


def _order_rpr(rpr) -> None:
    kids = list(rpr)
    kids.sort(key=lambda e: _RPR_ORDER.index(tag(e)) if tag(e) in _RPR_ORDER else len(_RPR_ORDER))
    for k in kids:
        rpr.append(k)


def fit_table(doc, tbl, total: int, keep_if_fits: bool = False, force: bool = False) -> str:
    """Give each column at least its widest unbreakable token and share the
    rest by content length, with fixed layout so Word and LibreOffice
    agree. When the minimums exceed ``total`` twips, the table font steps
    down to no less than 8pt. Returns ``"fit"``, ``"kept"`` (a template
    table whose own widths already hold its content) or ``"landscape"``
    (it cannot fit this width; the caller gives it a landscape section).
    With ``force`` it always fits, at 8pt if it must, and returns
    ``"squeezed"`` when even that leaves a column under its longest word."""
    base = table_font_pt(doc, tbl)
    mins, nats = column_needs(doc, tbl, base)
    if keep_if_fits:
        grid = tbl.find(qn("w:tblGrid"))
        have = [int(gc.get(qn("w:w")) or 0) for gc in grid.findall(qn("w:gridCol"))] if grid is not None else []
        if len(have) == len(mins) and all(h >= m for h, m in zip(have, mins)) and sum(have) <= total:
            return "kept"
    pt = base
    while sum(mins) > total and pt - 0.5 >= MIN_TABLE_PT:
        pt -= 0.5
        mins, nats = column_needs(doc, tbl, pt)
    if sum(mins) > total and not force:
        if pt != base:
            mins, nats = column_needs(doc, tbl, base)
        return "landscape"
    if pt != base:
        _set_table_pt(tbl, pt, base)
    squeezed = sum(mins) > total
    if squeezed:  # forced: scale the minimums down; nothing better is possible
        widths = [m * total / sum(mins) for m in mins]
    elif sum(nats) <= total:
        widths = [n + (total - sum(nats)) * n / sum(nats) for n in nats]
    else:
        room = total - sum(mins)
        want = [max(n - m, 0) for m, n in zip(mins, nats)]
        sw = sum(want) or 1
        widths = [m + room * w / sw for m, w in zip(mins, want)]
    ints = [int(-(-w // 1)) for w in widths]  # round up, so no column falls under its minimum
    slack = [i - m for i, m in zip(ints, mins)]
    ints[slack.index(max(slack))] -= sum(ints) - total  # the rounding excess comes from the roomiest column
    _set_widths(tbl, ints, total)
    return "squeezed" if squeezed else "fit"


def landscape_of(sect):
    """A landscape copy of a section's properties (same margins)."""
    land = copy.deepcopy(sect)
    pg = land.find(qn("w:pgSz"))
    w, h = int(pg.get(qn("w:w"))), int(pg.get(qn("w:h")))
    if w < h:
        pg.set(qn("w:w"), str(h))
        pg.set(qn("w:h"), str(w))
    pg.set(qn("w:orient"), "landscape")
    _plain_section(land)
    return land


def _plain_section(sect) -> None:
    for name in ("w:titlePg", "w:type"):
        for el in sect.findall(qn(name)):
            sect.remove(el)


_SECT_ORDER = ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
               "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
               "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange"]


def _order_sectpr(sect) -> None:
    kids = list(sect)
    kids.sort(key=lambda e: _SECT_ORDER.index(tag(e)) if tag(e) in _SECT_ORDER else len(_SECT_ORDER))
    for k in kids:
        sect.append(k)


def section_break(sect):
    """A paragraph that ends a section with ``sect``'s properties."""
    p = OxmlElement("w:p")
    p.get_or_add_pPr().append(copy.deepcopy(sect))
    return p


# ============================================================ table of contents

def _bookmark_ids(body) -> int:
    ids = [int(b.get(qn("w:id"))) for b in body.iter(qn("w:bookmarkStart"))
           if (b.get(qn("w:id")) or "").lstrip("-").isdigit()]
    return max(ids, default=0) + 1


def heading_bookmark(p_el, name: str, next_id: list[int]) -> str:
    """The heading's _Toc bookmark name, adding ``name`` when it has none."""
    for b in p_el.findall(qn("w:bookmarkStart")):
        if (b.get(qn("w:name")) or "").startswith("_Toc"):
            return b.get(qn("w:name"))
    bid = str(next_id[0])
    next_id[0] += 1
    bs = make("w:bookmarkStart", w_id=bid, w_name=name)
    be = make("w:bookmarkEnd", w_id=bid)
    ppr = p_el.find(qn("w:pPr"))
    if ppr is not None:
        ppr.addnext(bs)
    else:
        p_el.insert(0, bs)
    p_el.append(be)
    return name


def _toc_field(body):
    """(begin paragraph, begin run, end paragraph, end run) of the TOC
    field, or None."""
    for it in body.iter(qn("w:instrText")):
        if not (it.text or "").strip().upper().startswith("TOC"):
            continue
        r = it.getparent()
        p = r.getparent()
        while p is not None and tag(p) != "p":
            p = p.getparent()
        if p is None:
            continue
        # the begin run: the closest fldChar begin before the instruction in this paragraph
        runs = [x for x in p.iter(qn("w:r"))]
        idx = runs.index(r)
        begin = None
        for x in reversed(runs[:idx]):
            fc = x.find(qn("w:fldChar"))
            if fc is not None and fc.get(qn("w:fldCharType")) == "begin":
                begin = x
                break
        if begin is None:
            continue
        depth = 0
        cur = p
        started = False
        while cur is not None:
            if tag(cur) == "p":
                for x in cur.iter(qn("w:r")):
                    if not started:
                        if x is begin:
                            started = True
                        else:
                            continue
                    fc = x.find(qn("w:fldChar"))
                    if fc is None:
                        continue
                    t = fc.get(qn("w:fldCharType"))
                    if t == "begin":
                        depth += 1
                    elif t == "end":
                        depth -= 1
                        if depth == 0:
                            return p, begin, cur, x
            cur = cur.getnext()
        return None
    return None


def rebuild_toc(doc, entries: list[tuple[int, str | None, str, str]], toc_styles: dict[int, str]) -> int:
    """Rebuild the TOC field's cached entries. ``entries`` are (level,
    number or None, title, bookmark). Page numbers are carried over from
    the old cache where the bookmark is the same; the field is marked dirty
    and ``updateFields`` set, so Word repaginates it on open. Returns the
    entry count, or raises DocxBuildError when there is no TOC field."""
    body = doc.element.body
    found = _toc_field(body)
    if found is None:
        raise DocxBuildError("the template has no table of contents field (a TOC field) to rebuild")
    begin_p, begin_r, end_p, end_r = found
    old_pages: dict[str, str] = {}
    protos: dict[str, Any] = {}
    cur = begin_p
    span = []
    while cur is not None:
        span.append(cur)
        if cur is end_p:
            break
        cur = cur.getnext()
    for p in span:
        if tag(p) != "p":
            continue
        sid = pstyle_id(p)
        if sid and sid not in protos and p.find(qn("w:pPr")) is not None:
            protos[sid] = copy.deepcopy(p.find(qn("w:pPr")))
        for hl in p.iter(qn("w:hyperlink")):
            ts = [t.text for t in hl.iter(qn("w:t")) if t.text]
            if ts and ts[-1].strip().isdigit():
                old_pages[hl.get(qn("w:anchor"))] = ts[-1].strip()
    # the field's opening runs: begin through separate, as written in the begin paragraph
    runs = [x for x in begin_p if tag(x) == "r"]
    opening = []
    on = False
    for x in begin_p.iter(qn("w:r")):
        if x is begin_r:
            on = True
        if on:
            opening.append(copy.deepcopy(x))
            fc = x.find(qn("w:fldChar"))
            if fc is not None and fc.get(qn("w:fldCharType")) == "separate":
                break
    del runs
    for fc in opening[0].iter(qn("w:fldChar")):
        fc.set(qn("w:dirty"), "true")
    lead_text = []
    for x in begin_p:
        if x is begin_r or begin_r in list(x.iter(qn("w:r"))):
            break
        if tag(x) == "r" and ptext(x).strip():
            lead_text.append(copy.deepcopy(x))
    end_alone = end_p is not begin_p and not ptext(end_p).strip()
    new_ps = []
    for i, (level, num, text, bm) in enumerate(entries):
        sid = toc_styles.get(level) or toc_styles.get(1)
        p = OxmlElement("w:p")
        ppr = copy.deepcopy(protos[sid]) if sid in protos else None
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            ps = OxmlElement("w:pStyle")
            ps.set(qn("w:val"), sid)
            ppr.append(ps)
        for rpr in ppr.findall(qn("w:rPr")):
            ppr.remove(rpr)
        p.append(ppr)
        if i == 0:
            for x in opening:
                p.append(x)
        hl = make("w:hyperlink", w_anchor=bm, w_history="1")
        if num:
            hl.append(_text_run(num))
            hl.append(_tab_run())
        hl.append(_text_run(text))
        hl.append(_tab_run())
        page = old_pages.get(bm, "")
        hl.append(_fld_run("begin"))
        hl.append(_instr_run(f" PAGEREF {bm} \\h "))
        hl.append(_fld_run("separate"))
        hl.append(_text_run(page))
        hl.append(_fld_run("end"))
        p.append(hl)
        new_ps.append(p)
    if not new_ps:
        raise DocxBuildError("no headings to put in the table of contents")
    tail = None
    if end_alone:
        tail = OxmlElement("w:p")
        if end_p.find(qn("w:pPr")) is not None:
            tail.append(copy.deepcopy(end_p.find(qn("w:pPr"))))
        tail.append(copy.deepcopy(end_r))
    else:
        new_ps[-1].append(copy.deepcopy(end_r))
    if lead_text:
        lp = OxmlElement("w:p")
        if begin_p.find(qn("w:pPr")) is not None:
            lp.append(copy.deepcopy(begin_p.find(qn("w:pPr"))))
        for x in lead_text:
            lp.append(x)
        new_ps.insert(0, lp)
    for p in new_ps:
        begin_p.addprevious(p)
    if tail is not None:
        begin_p.addprevious(tail)
    for p in span:
        if has_sect(p):
            neutral = OxmlElement("w:p")
            neutral.append(copy.deepcopy(p.find(qn("w:pPr"))))
            p.addprevious(neutral)
        p.getparent().remove(p)
    set_update_fields(doc)
    return len(entries)


def _text_run(text: str):
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _tab_run():
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:tab"))
    return r


def _fld_run(kind: str):
    r = OxmlElement("w:r")
    r.append(make("w:fldChar", w_fldCharType=kind))
    return r


def _instr_run(text: str):
    r = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = text
    r.append(it)
    return r


def set_update_fields(doc) -> None:
    settings = doc.settings.element
    uf = settings.find(qn("w:updateFields"))
    if uf is None:
        uf = OxmlElement("w:updateFields")
        # the schema orders settings: updateFields comes before these
        after = {"hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                 "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                 "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType",
                 "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"}
        anchor = next((ch for ch in settings if tag(ch) in after), None)
        if anchor is not None:
            anchor.addprevious(uf)
        else:
            settings.append(uf)
    uf.set(qn("w:val"), "true")


# ============================================================ reading back

@dataclass
class Structure:
    """What a .docx holds, read back: headings with rendered numbers, the
    TOC's cached entries, tables, pictures, captions and section count."""

    headings: list[tuple[int, str | None, str]] = field(default_factory=list)
    toc: list[tuple[int, str | None, str]] = field(default_factory=list)
    tables: list[list[list[str]]] = field(default_factory=list)
    pictures: list[tuple[int, int]] = field(default_factory=list)  # (cx, cy) EMU
    captions: list[str] = field(default_factory=list)
    texts: list[str] = field(default_factory=list)
    list_items: list[tuple[int, str | None, str]] = field(default_factory=list)  # (ilvl, numFmt, text)
    sections: int = 0
    update_fields: bool = False
    toc_dirty: bool = False
    revisions: int = 0
    comments: int = 0
    styles: dict[str, str] = field(default_factory=dict)
    narrow_tables: list[int] = field(default_factory=list)  # tables with a column under its longest token

    def all_text(self) -> str:
        return "\n".join(self.texts + [c for t in self.tables for r in t for c in r])


def _format(n: int, fmt: str) -> str:
    if fmt in ("upperLetter", "lowerLetter"):
        s = ""
        while n > 0:
            n, r = divmod(n - 1, 26)
            s = chr(65 + r) + s
        return s if fmt == "upperLetter" else s.lower()
    if fmt in ("upperRoman", "lowerRoman"):
        vals = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"), (50, "L"),
                (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
        s = ""
        for v, sym in vals:
            while n >= v:
                s += sym
                n -= v
        return s if fmt == "upperRoman" else s.lower()
    return str(n)


class NumberingSim:
    """Simulates Word list numbering to give each paragraph its rendered
    label."""

    def __init__(self, doc):
        self.nums: dict[str, tuple[str, dict[int, int]]] = {}
        self.abstracts: dict[str, dict[int, dict]] = {}
        try:
            el = doc.part.numbering_part.element
        except Exception:  # pragma: no cover - no numbering part
            el = None
        if el is None:
            return
        for a in el.findall(qn("w:abstractNum")):
            levels = {}
            for lvl in a.findall(qn("w:lvl")):
                il = int(lvl.get(qn("w:ilvl")))

                def val(name, default=None, lvl=lvl):
                    x = lvl.find(qn(name))
                    return x.get(qn("w:val")) if x is not None else default
                levels[il] = {"start": int(val("w:start", "1")), "fmt": val("w:numFmt", "decimal"),
                              "text": val("w:lvlText", ""), "restart": val("w:lvlRestart")}
            self.abstracts[a.get(qn("w:abstractNumId"))] = levels
        for n in el.findall(qn("w:num")):
            starts = {}
            for ov in n.findall(qn("w:lvlOverride")):
                so = ov.find(qn("w:startOverride"))
                if so is not None:
                    starts[int(ov.get(qn("w:ilvl")))] = int(so.get(qn("w:val")))
            self.nums[n.get(qn("w:numId"))] = (n.find(qn("w:abstractNumId")).get(qn("w:val")), starts)
        self.counters: dict[str, list[int | None]] = {}
        self.seen_num: set[str] = set()

    def label(self, num_id: str | None, ilvl: int) -> tuple[str | None, str | None]:
        """(label, numFmt) for the next paragraph on num_id at ilvl."""
        if not num_id or num_id == "0" or num_id not in self.nums:
            return None, None
        aid, starts = self.nums[num_id]
        levels = self.abstracts.get(aid, {})
        lv = levels.get(ilvl)
        if lv is None:
            return None, None
        key = aid
        cnt = self.counters.setdefault(key, [None] * 9)
        if num_id not in self.seen_num:
            self.seen_num.add(num_id)
            for k, s in starts.items():
                cnt[k] = None
        if cnt[ilvl] is None:
            cnt[ilvl] = starts.get(ilvl, lv["start"])
        else:
            cnt[ilvl] += 1
        for k in range(ilvl + 1, 9):
            cnt[k] = None
        if lv["fmt"] in ("bullet", "none"):
            return lv["text"], lv["fmt"]

        def sub(m):
            k = int(m.group(1)) - 1
            v = cnt[k] if cnt[k] is not None else levels.get(k, {}).get("start", 1)
            return _format(v, levels.get(k, {}).get("fmt", "decimal"))
        return re.sub(r"%(\d)", sub, lv["text"]).rstrip("."), lv["fmt"]


def paragraph_num(p, styles: dict[str, StyleInfo]) -> tuple[str | None, int]:
    ppr = p.find(qn("w:pPr"))
    sid = pstyle_id(p)
    num_id, ilvl = None, None
    if sid and sid in styles and styles[sid].num:
        num_id, ilvl = styles[sid].num
    if ppr is not None and ppr.find(qn("w:numPr")) is not None:
        np_ = ppr.find(qn("w:numPr"))
        n = np_.find(qn("w:numId"))
        il = np_.find(qn("w:ilvl"))
        if n is not None:
            num_id = n.get(qn("w:val"))
        if il is not None:
            ilvl = il.get(qn("w:val"))
    return num_id, int(ilvl or 0)


def _toc_span(body) -> list:
    found = _toc_field(body)
    if not found:
        return []
    out, cur = [], found[0]
    while cur is not None:
        out.append(cur)
        if cur is found[2]:
            break
        cur = cur.getnext()
    return out


def walk(doc):
    """Every paragraph in reading order as ``(p, where, label, fmt, ilvl,
    outline)``: ``where`` is ``toc`` (inside the TOC field), ``box`` (a
    text box), ``table`` or ``flow``; ``label`` and ``fmt`` are the rendered
    list label and its number format, from simulating Word's numbering
    (text boxes and the TOC do not advance it)."""
    body = doc.element.body
    styles = style_table(doc)
    sim = NumberingSim(doc)
    toc = set(_toc_span(body))
    for p in body.iter(qn("w:p")):
        where = "flow"
        anc = p.getparent()
        while anc is not None and anc is not body:
            if anc in toc:
                where = "toc"
                break
            if tag(anc) == "txbxContent":
                where = "box"
                break
            if tag(anc) == "tbl" and where == "flow":
                where = "table"
            anc = anc.getparent()
        if p in toc:
            where = "toc"
        num_id, ilvl = paragraph_num(p, styles)
        label, fmt = (None, None) if where in ("toc", "box") else sim.label(num_id, ilvl)
        yield p, where, label, fmt, ilvl, paragraph_outline(p, styles)


def heading_paragraphs(doc):
    """``(p, level, number, fmt, text)`` for each outline-level 1 or 2
    paragraph with text in the document flow; ``number`` is its rendered
    number, or None when it is unnumbered."""
    for p, where, label, fmt, _ilvl, outline in walk(doc):
        if where != "flow" or outline not in (0, 1):
            continue
        text = own_text(p).strip()
        if text:
            yield p, outline + 1, (label if fmt not in (None, "bullet", "none") else None), fmt, text


def narrow_tables(doc) -> list[int]:
    """Indexes of body tables where a column is narrower than the widest
    unbreakable token in it, at the size the table's runs are set to."""
    bad = []
    for i, tbl in enumerate(doc.element.body.iter(qn("w:tbl"))):
        grid = tbl.find(qn("w:tblGrid"))
        have = [int(gc.get(qn("w:w")) or 0) for gc in grid.findall(qn("w:gridCol"))] if grid is not None else []
        # at the table's own size: a run that carries an explicit size is measured at that size
        mins, _ = column_needs(doc, tbl, table_font_pt(doc, tbl))
        # widths are whole twips: rounding a table that only just fits can cost a column a twip or two
        if len(have) == len(mins) and any(h + len(have) < m for h, m in zip(have, mins)):
            bad.append(i)
    return bad


def read_structure(path: str | Path) -> Structure:
    """Read a .docx back: rendered heading numbers (simulating Word's list
    numbering), the TOC's cached entries, tables, pictures and captions."""
    doc = Document(str(path))
    body = doc.element.body
    styles = style_table(doc)
    out = Structure()
    caption_ids = {sid for sid, s in styles.items() if s.name.lower() == "caption"}
    out.styles = {sid: s.name for sid, s in styles.items()}
    toc_paras = _toc_span(body)
    found = _toc_field(body)
    if found:
        for fc in found[1].iter(qn("w:fldChar")):
            out.toc_dirty = fc.get(qn("w:dirty")) in ("true", "1")
        for p in toc_paras:
            if tag(p) != "p":
                continue
            for hl in p.iter(qn("w:hyperlink")):
                parts = []
                cur_t = []
                for r in hl.iter(qn("w:r")):
                    if r.find(qn("w:tab")) is not None:
                        parts.append("".join(cur_t))
                        cur_t = []
                    for t in r.findall(qn("w:t")):
                        cur_t.append(t.text or "")
                parts.append("".join(cur_t))
                sid = pstyle_id(p) or ""
                lvl = int(re.sub(r"\D", "", styles[sid].name) or 1) if sid in styles else 1
                if len(parts) >= 3:
                    out.toc.append((lvl, parts[0], parts[1]))
                elif parts:
                    out.toc.append((lvl, None, parts[0]))
    for p, where, label, fmt, ilvl, lvl in walk(doc):
        if where != "flow":
            if where == "box" and own_text(p).strip():
                out.texts.append(own_text(p))
            continue
        text = own_text(p)
        if lvl is not None and text.strip():
            out.headings.append((lvl + 1, label if fmt not in (None, "bullet", "none") else None, text.strip()))
        elif text.strip():
            out.texts.append(text)
            if label is not None:
                out.list_items.append((ilvl, fmt, text))
            if pstyle_id(p) in caption_ids:
                out.captions.append(text)
    for t in body.iter(qn("w:tbl")):
        rows = []
        for tr in t.findall(qn("w:tr")):
            rows.append([ptext(tc) for tc in tr.findall(qn("w:tc"))])
        out.tables.append(rows)
    out.narrow_tables = narrow_tables(doc)
    for ext in body.iter(qn("wp:extent")):
        out.pictures.append((int(ext.get("cx")), int(ext.get("cy"))))
    out.sections = len(list(body.iter(qn("w:sectPr"))))
    uf = doc.settings.element.find(qn("w:updateFields"))
    out.update_fields = uf is not None and uf.get(qn("w:val")) in ("true", "1")
    out.revisions = sum(1 for n in ("w:ins", "w:del", "w:moveFrom", "w:moveTo", "w:rPrChange", "w:pPrChange")
                        for _ in body.iter(qn(n)))
    out.comments = sum(1 for _ in body.iter(qn("w:commentReference"))) + sum(
        1 for rel in doc.part.rels.values() if "comments" in rel.reltype)
    return out


# ============================================================ shared build steps

def output_path(project, out_dir: Path) -> Path:
    name = re.sub(r'[\\/:*?"<>|]+', "-", f"{project.title} v{project.version}.docx")
    return out_dir / name


def template_path(project) -> Path | None:
    """The template to fill: arch-doc.yaml ``template`` (relative to the
    project), else the profile's ``template.path`` (relative to the
    profile file). None for the neutral document."""
    t = project.config.get("template")
    if t:
        path = Path(t)
        if not path.is_absolute():
            path = project.root / path
    else:
        tp = project.profile.get("template.path")
        if not tp:
            return None
        path = Path(tp)
        if not path.is_absolute():
            base = project.profile.path.parent if project.profile.path else project.root
            path = base / path
    if not path.exists():
        raise DocxBuildError(f"template {path} does not exist (arch-doc.yaml template / profile template.path)")
    return path


def _person(d: Any) -> tuple[str, str]:
    """A person's name and position (role and team, less any the name
    already says, and less a bare TBC when something else is known)."""
    if isinstance(d, dict):
        name, details = person(d)
        pos = ", ".join(x for x in details if x != "TBC") or ", ".join(details)
        return name, pos
    return (str(d) if d else ""), ""


def _date_text(v: Any) -> str:
    if isinstance(v, (_dt.date, _dt.datetime)):
        return v.isoformat()
    return "" if v is None else str(v)


def history_rows(project) -> list[dict[str, str]]:
    fmt = date_format(project)
    return [{"date": format_date(h.get("date"), fmt), "version": str(h.get("version", "")),
             "author": str(h.get("author", "")), "change": str(h.get("change", ""))}
            for h in project.config.get("history") or [] if isinstance(h, dict)]


def heading_entries(project) -> list[tuple[int, str, str, str]]:
    """(level, number, title, anchor) for every chapter and section."""
    out = []
    for ch in project.chapters:
        out.append((1, str(ch.number), ch.title, ch.anchor))
        for s in ch.sections:
            out.append((2, s.number, s.title, s.anchor))
    return out


def bookmark_name(anchor: str) -> str:
    return "_Toc_" + re.sub(r"[^A-Za-z0-9]", "_", anchor)[:34]


def build(project, out_dir: str | Path | None = None) -> list[Path]:
    """Build ``<title> v<version>.docx`` in ``out_dir`` (default
    ``<project>/out``): the template filled in place when the project or
    profile names one, the neutral document otherwise."""
    out = Path(out_dir) if out_dir is not None else project.root / "out"
    out.mkdir(parents=True, exist_ok=True)
    tpl = template_path(project)
    target = output_path(project, out)
    if tpl is not None:
        TemplateFill(project, tpl).build(target)
    else:
        NeutralBuild(project).build(target)
    return [target]


def _verify_numbering(project, path: Path, unnumbered_ok: set[str] = frozenset()) -> None:
    """The rendered heading numbers must be the core's numbers."""
    st = read_structure(path)
    expected = [(lvl, num, title) for lvl, num, title, _ in heading_entries(project)]
    got = [(lvl, num, title) for lvl, num, title in st.headings if lvl in (1, 2)]
    if expected:
        # headings ahead of the first chapter (front matter) are not chapters
        first = next((i for i, (lvl, _, t) in enumerate(got) if lvl == 1 and t == expected[0][2]), 0)
        got = got[first:]
    if [(lvl, t) for lvl, _, t in got] != [(lvl, t) for lvl, _, t in expected]:
        a = [t for _, _, t in got]
        b = [t for _, _, t in expected]
        diff = "\n".join(difflib.unified_diff(b, a, "metamodel", "docx", lineterm="", n=1))
        raise DocxBuildError(f"{path.name}: the headings do not match the document's chapters and sections:\n{diff}")
    for (lvl, num, title), (_, want, _) in zip(got, expected):
        if num != want and not (num is None and title in unnumbered_ok):
            raise DocxBuildError(f"{path.name}: '{title}' renders as {num or 'unnumbered'}, "
                                 f"but the document numbers it {want}")


# ============================================================ neutral document

class NeutralBuild:
    """The neutral document: cover, document control, TOC, numbered
    chapters, house styles."""

    def __init__(self, project):
        self.p = project
        self.doc = Document()
        self.palette = project.profile.palette
        self.fonts = project.profile.fonts

    def build(self, target: Path) -> None:
        d = self.doc
        self._page()
        self._styles()
        num = Numbering(d)
        heading_num = num.heading_definition([self.sid("Heading 1"), self.sid("Heading 2")])
        for name, lvl in (("Heading 1", 0), ("Heading 2", 1)):
            ppr = d.styles[name].element.get_or_add_pPr()
            np_ = ppr.get_or_add_numPr()
            np_.get_or_add_ilvl().val = lvl
            np_.get_or_add_numId().val = int(heading_num)
        styles = Styles(self.sid("Heading 1"), self.sid("Heading 2"), self.sid("Normal"),
                        self.sid("List Paragraph"), self.sid("House Table"), self.sid("Caption"),
                        self.sid("toc 1"), self.sid("toc 2"))
        bullet = num.bullet_definition(self.fonts.get("body") or "Calibri")
        ordered = num.ordered_abstract()
        self.r = Renderer(self.p, d, styles, num, bullet, ordered)
        body = d.element.body
        for el in list(body):
            if tag(el) != "sectPr":
                body.remove(el)
        self.sect = body.find(qn("w:sectPr"))
        self._footer()
        self._cover()
        self._control()
        self._toc_placeholder()
        next_id = [_bookmark_ids(body)]
        entries = []
        for ch in self.p.chapters:
            h = self.r.para(styles.h1)
            h.add_run(ch.title)
            self._append(h._p)
            entries.append((1, str(ch.number), ch.title, heading_bookmark(h._p, bookmark_name(ch.anchor), next_id)))
            for el in self.r.blocks(ch.lead, self.sect):
                self._append(el)
            for s in ch.sections:
                h2 = self.r.para(styles.h2)
                h2.add_run(s.title)
                self._append(h2._p)
                entries.append((2, s.number, s.title, heading_bookmark(h2._p, bookmark_name(s.anchor), next_id)))
                for el in self.r.blocks(s.blocks, self.sect):
                    self._append(el)
        rebuild_toc(d, entries, {1: styles.toc1, 2: styles.toc2})
        cp = d.core_properties
        cp.title = self.p.title
        cp.subject = f"{self.p.kind.capitalize()} design"
        cp.author = _person(self.p.config.get("author"))[0]
        cp.comments = ""
        cp.revision = 1
        d.save(str(target))
        _verify_numbering(self.p, target)
        for w in self.r.squeezed:
            print(f"archdoc build docx: warning: {w}", file=sys.stderr)

    # ------------------------------------------------------------ helpers
    def sid(self, name: str) -> str:
        return self.doc.styles[name].style_id

    def _append(self, el) -> None:
        self.sect.addprevious(el)

    def _extend(self, els) -> None:
        for el in els:
            self._append(el)

    def _page(self) -> None:
        s = self.doc.sections[0]
        s.page_width, s.page_height = Twips(11906), Twips(16838)  # A4
        for side in ("left_margin", "right_margin"):
            setattr(s, side, Twips(1247))
        s.top_margin, s.bottom_margin = Twips(1304), Twips(1247)
        s.header_distance, s.footer_distance = Twips(567), Twips(567)

    def _font(self, style, family: str | None, size: float | None = None, colour: str | None = None,
              bold: bool | None = None, italic: bool | None = None) -> None:
        f = style.font
        if family:
            f.name = family
            rf = style.element.get_or_add_rPr().get_or_add_rFonts()
            for a in ("w:eastAsia", "w:cs"):
                rf.set(qn(a), family)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
        if size:
            f.size = Pt(size)
        if colour:
            f.color.rgb = RGBColor.from_string(_hex(colour))
            c = style.element.rPr.find(qn("w:color"))
            for a in ("w:themeColor", "w:themeShade", "w:themeTint"):
                if c is not None and c.get(qn(a)) is not None:
                    del c.attrib[qn(a)]
        if bold is not None:
            f.bold = bold
        if italic is not None:
            f.italic = italic

    def _styles(self) -> None:
        d, pal, fonts = self.doc, self.palette, self.fonts
        body_font, head_font = fonts.get("body"), fonts.get("heading")
        normal = d.styles["Normal"]
        self._font(normal, body_font, 10.5, pal["ink"])
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for name, size, colour, before, after in (("Heading 1", 18, pal["ink"], 0, 10),
                                                   ("Heading 2", 13, pal["accent"], 14, 6)):
            st = d.styles[name]
            self._font(st, head_font, size, colour, bold=True)
            pf = st.paragraph_format
            pf.space_before, pf.space_after = Pt(before), Pt(after)
            pf.keep_with_next = True
            if name == "Heading 1":
                pf.page_break_before = True
        title = d.styles["Title"]
        self._font(title, head_font, 30, pal["ink"])
        tppr = title.element.get_or_add_pPr()
        for b in tppr.findall(qn("w:pBdr")):
            tppr.remove(b)
        title.paragraph_format.space_after = Pt(12)
        sub = d.styles["Subtitle"]
        self._font(sub, body_font, 14, pal["muted"], italic=False)
        cap = d.styles["Caption"]
        self._font(cap, body_font, 9, pal["muted"], bold=False, italic=True)
        cap.paragraph_format.space_after = Pt(12)
        lp = d.styles["List Paragraph"]
        self._font(lp, body_font)
        lp.paragraph_format.space_after = Pt(3)
        front = d.styles.add_style("Front Heading", WD_STYLE_TYPE.PARAGRAPH)
        front.base_style = normal
        self._font(front, head_font, 16, pal["ink"], bold=True)
        front.paragraph_format.space_before = Pt(0)
        front.paragraph_format.space_after = Pt(10)
        front.paragraph_format.keep_with_next = True
        label = d.styles.add_style("Front Label", WD_STYLE_TYPE.PARAGRAPH)
        label.base_style = normal
        self._font(label, body_font, 10.5, pal["accent"], bold=True)
        label.paragraph_format.space_before = Pt(10)
        label.paragraph_format.keep_with_next = True
        tw = text_box(self.doc.element.body.find(qn("w:sectPr")))[0] // EMU_PER_TWIP
        for i, (size, bold, indent) in enumerate(((11, True, 0), (10.5, False, 284)), start=1):
            st = d.styles.add_style(f"toc {i}", WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = normal
            st.next_paragraph_style = normal
            self._font(st, body_font, size, pal["ink"], bold=bold)
            st.paragraph_format.space_before = Pt(6 if i == 1 else 0)
            st.paragraph_format.space_after = Pt(2)
            ppr = st.element.get_or_add_pPr()
            tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="left" w:pos="{indent + 567}"/>'
                             f'<w:tab w:val="right" w:leader="dot" w:pos="{tw}"/></w:tabs>')
            ppr.insert_element_before(tabs, "w:spacing", "w:ind", "w:jc", "w:outlineLvl", "w:rPr")
            if indent:
                st.paragraph_format.left_indent = Twips(indent)
            st.element.set(qn("w:customStyle"), "0")
        self._table_style()
        self._character_styles()

    def _table_style(self) -> None:
        pal = self.palette
        rule, accent, panel = _hex(pal["rule"]), _hex(pal["accent"]), _hex(pal["panel"])
        border = "".join(f'<w:{s} w:val="single" w:sz="4" w:space="0" w:color="{rule}"/>'
                         for s in ("top", "left", "bottom", "right", "insideH", "insideV"))
        xml = (f'<w:style {nsdecls("w")} w:type="table" w:customStyle="1" w:styleId="HouseTable">'
               f'<w:name w:val="House Table"/><w:basedOn w:val="TableNormal"/><w:uiPriority w:val="40"/>'
               f'<w:pPr><w:spacing w:before="30" w:after="30" w:line="240" w:lineRule="auto"/></w:pPr>'
               f'<w:rPr><w:sz w:val="18"/><w:szCs w:val="18"/></w:rPr>'
               f'<w:tblPr><w:tblBorders>{border}</w:tblBorders>'
               f'<w:tblCellMar><w:top w:w="40" w:type="dxa"/><w:left w:w="85" w:type="dxa"/>'
               f'<w:bottom w:w="40" w:type="dxa"/><w:right w:w="85" w:type="dxa"/></w:tblCellMar></w:tblPr>'
               f'<w:tblStylePr w:type="firstRow"><w:rPr><w:b/><w:bCs/><w:color w:val="FFFFFF"/></w:rPr>'
               f'<w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="{accent}"/></w:tcPr></w:tblStylePr>'
               f'<w:tblStylePr w:type="band1Horz"><w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="{panel}"/>'
               f'</w:tcPr></w:tblStylePr></w:style>')
        self.doc.styles.element.append(parse_xml(xml))

    def _character_styles(self) -> None:
        st = self.doc.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
        st.font.color.rgb = RGBColor.from_string(_hex(self.palette["accent"]))
        st.font.underline = True

    def _cover(self) -> None:
        p, r = self.p, self.r
        for _ in range(6):
            self._append(r.para(self.sid("Normal"))._p)
        org = p.profile.get("org.name")
        if org:
            lab = r.para(self.sid("Front Label"))
            lab.add_run(str(org))
            self._append(lab._p)
        t = r.para(self.sid("Title"))
        t.add_run(p.title)
        self._append(t._p)
        sub = r.para(self.sid("Subtitle"))
        sub.add_run(f"{p.kind.capitalize()} design")
        self._append(sub._p)
        meta = r.para(self.sid("Normal"))
        meta.paragraph_format.space_before = Pt(24)
        bits = [f"Version {p.version}"]
        if p.config.get("date"):
            bits.append(format_date(p.config["date"], date_format(p)))
        meta.add_run(" · ".join(bits)).bold = True
        self._append(meta._p)
        for label, who in (("Owner", p.config.get("owner")), ("Author", p.config.get("author"))):
            name, pos = _person(who)
            if not name:
                continue
            line = r.para(self.sid("Normal"))
            line.paragraph_format.space_after = Pt(0)
            line.add_run(f"{label}: ").bold = True
            line.add_run(name + (f", {pos}" if pos else ""))
            self._append(line._p)
        deps = [d for d in p.config.get("depends_on") or [] if isinstance(d, dict)]
        if deps:
            line = r.para(self.sid("Normal"))
            line.paragraph_format.space_before = Pt(12)
            line.add_run("Depends on: ").bold = True
            line.add_run("; ".join(f"{d.get('title')} v{d.get('version')}" if d.get("version")
                                   else str(d.get("title")) for d in deps))
            self._append(line._p)
        # the cover is its own section, a title page with no footer
        cover = copy.deepcopy(self.sect)
        if cover.find(qn("w:titlePg")) is None:
            cover.append(OxmlElement("w:titlePg"))
        _order_sectpr(cover)
        self._append(section_break(cover))

    def _control(self) -> None:
        p, r = self.p, self.r
        h = r.para(self.sid("Front Heading"))
        h.add_run("Document control")
        self._append(h._p)
        lab = r.para(self.sid("Front Label"))
        lab.add_run("Ownership")
        self._append(lab._p)
        rows = []
        for label, who in (("Owner", p.config.get("owner")), ("Author", p.config.get("author"))):
            name, pos = _person(who)
            rows.append([label, name, pos])
        self._extend(r.table(["Role", "Name", "Position"], rows, self.sect))
        lab = r.para(self.sid("Front Label"))
        lab.add_run("Version history")
        self._append(lab._p)
        hist = [[h["date"], h["version"], h["author"], h["change"]] for h in history_rows(p)]
        self._extend(r.table(["Date", "Version", "Author", "Change"], hist, self.sect, empty_note="No releases yet."))
        deps = [d for d in p.config.get("depends_on") or [] if isinstance(d, dict)]
        if deps:
            lab = r.para(self.sid("Front Label"))
            lab.add_run("Designs this one depends on")
            self._append(lab._p)
            self._extend(r.table(["Design", "Kind", "Version", "Reference"],
                                 [[str(d.get("title", "")), str(d.get("kind", "")), str(d.get("version", "")),
                                   str(d.get("ref", ""))] for d in deps], self.sect))

    def _toc_placeholder(self) -> None:
        r = self.r
        brk = r.para(self.sid("Normal"))
        brk.add_run().add_break(WD_BREAK.PAGE)
        self._append(brk._p)
        h = r.para(self.sid("Front Heading"))
        h.add_run("Contents")
        self._append(h._p)
        sdt = parse_xml(f'<w:sdt {nsdecls("w")}><w:sdtPr><w:docPartObj>'
                        f'<w:docPartGallery w:val="Table of Contents"/><w:docPartUnique/></w:docPartObj>'
                        f'</w:sdtPr><w:sdtContent/></w:sdt>')
        content = sdt.find(qn("w:sdtContent"))
        first = r.para(self.sid("toc 1"))
        first._p.append(_fld_run("begin"))
        first._p.append(_instr_run(' TOC \\o "1-2" \\h \\z \\u '))
        first._p.append(_fld_run("separate"))
        last = r.para(self.sid("Normal"))
        last._p.append(_fld_run("end"))
        content.append(first._p)
        content.append(last._p)
        self._append(sdt)

    def _footer(self) -> None:
        s = self.doc.sections[0]
        f = s.footer
        par = f.paragraphs[0]
        par.style = self.doc.styles["Footer"]
        self._font(self.doc.styles["Footer"], self.fonts.get("body"), 8.5, self.palette["muted"])
        tw = text_box(self.sect)[0] // EMU_PER_TWIP
        ppr = par._p.get_or_add_pPr()
        for t in ppr.findall(qn("w:tabs")):
            ppr.remove(t)
        ppr.insert_element_before(parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="clear" w:pos="4680"/>'
                                            f'<w:tab w:val="clear" w:pos="9360"/>'
                                            f'<w:tab w:val="right" w:pos="{tw}"/></w:tabs>'),
                                  "w:spacing", "w:ind", "w:jc", "w:rPr")
        par.add_run(f"{self.p.short} v{self.p.version}")
        par.add_run("\t")
        self.r._field(par, " PAGE ", "1")


# ============================================================ template fill

@dataclass
class THeading:
    el: Any
    level: int
    title: str
    children: list["THeading"] = field(default_factory=list)


# Document-control stakeholder headers, by meaning, to stakeholder fields.
STAKEHOLDER_SYNONYMS = [
    (("business unit", "team", "department", "directorate", "division", "organisation", "organization",
      "unit"), "team"),
    (("interest", "concern"), "interest"),
    (("involvement", "raci", "participation"), "involvement"),
    (("name",), "name"),
    (("role", "stakeholder", "position", "title"), "title"),
]


class TemplateFill:
    """Fills an organisation's template in place."""

    def __init__(self, project, template: Path):
        self.p = project
        self.template = template
        self.doc = Document(str(template))
        self.body = self.doc.element.body
        self.warnings: list[str] = []

    # ------------------------------------------------------------ setup
    def build(self, target: Path) -> None:
        d = self.doc
        sections_before = len(list(self.body.iter(qn("w:sectPr"))))
        clean_revisions_and_comments(d)
        self.st_table = style_table(d)
        self.styles = self._resolve_styles()
        self.num = Numbering(d)
        self.bullet_num = self._template_bullet_num()
        self.ordered_abstract = self.num.ordered_abstract()
        self.r = TemplateRenderer(self.p, d, self.styles, self.num, self.bullet_num,
                                  self.ordered_abstract)
        self.direct_num = {lvl: self._direct_numpr(lvl) for lvl in (0, 1)}
        chapters = self._map_template()
        self._fill(chapters)
        self._neutralise_empty_headings()
        self._front_matter()
        self._fields()
        self._clean_bookmarks()
        entries, unnumbered = self._toc()
        rebuild_toc(d, entries, {1: self.styles.toc1, 2: self.styles.toc2})
        cp = d.core_properties
        cp.title = self.p.title
        cp.subject = f"{self.p.kind.capitalize()} design"
        cp.author = _person(self.p.config.get("author"))[0]
        cp.comments = ""
        d.save(str(target))
        sections_after = len(list(Document(str(target)).element.body.iter(qn("w:sectPr"))))
        if sections_after != sections_before:
            raise DocxBuildError(f"{target.name}: the template has {sections_before} sections but the "
                                 f"filled document has {sections_after}")
        _verify_numbering(self.p, target, unnumbered)
        self._warn_unnumbered_refs(unnumbered)
        for w in self.warnings + self.r.squeezed:
            print(f"archdoc build docx: warning: {w}", file=sys.stderr)

    def _resolve_styles(self) -> Styles:
        d = self.doc
        names = {s.name: sid for sid, s in self.st_table.items()}
        names_lower = {k.lower(): v for k, v in names.items()}
        # explicit profile mapping first, then what the template itself uses, then the profile defaults
        explicit = dict(((self.p.profile.raw.get("template") or {}).get("styles")) or {})
        defaults = dict(self.p.profile.get("template.styles") or {})
        detected = detect_styles(d)
        mapping = explicit

        def pick(role: str, required: bool = True) -> str | None:
            for cand in (explicit.get(role), detected.get(role), defaults.get(role)):
                if cand and cand in names:
                    return names[cand]
                if cand and cand.lower() in names_lower:
                    return names_lower[cand.lower()]
            if not required:
                return None
            raise DocxBuildError(
                f"template {self.template.name}: no '{role}' style: the profile's template.styles.{role} is "
                f"'{mapping.get(role)}', which the template does not define, and none could be detected "
                f"(run archdoc extract on the template for the mapping)")
        h1, h2, body = pick("h1"), pick("h2"), pick("body")
        bullet = pick("bullet", False) or self._ensure_style("List Paragraph", body)
        caption = pick("caption", False) or self._ensure_style("Caption", body, italic=True)
        toc1 = pick("toc1", False) or self._ensure_style("toc 1", body)
        toc2 = pick("toc2", False) or self._ensure_style("toc 2", body)
        table = None
        tname = explicit.get("table") or detected.get("table") or defaults.get("table")
        for cand in (explicit.get("table"), detected.get("table"), defaults.get("table")):
            sid = self._table_style_id(cand)
            if sid:
                table = sid
                break
        if table is None and tname:
            self.warnings.append(f"table style '{tname}' is not in the template; new tables are unstyled")
        return Styles(h1, h2, body, bullet, table, caption, toc1, toc2)

    def _table_style_id(self, name: str | None) -> str | None:
        if not name:
            return None
        for st in self.doc.styles.element.findall(qn("w:style")):
            if st.get(qn("w:type")) != "table":
                continue
            n = st.find(qn("w:name"))
            if (n is not None and n.get(qn("w:val")).lower() == name.lower()) or st.get(qn("w:styleId")) == name:
                return st.get(qn("w:styleId"))
        return None

    def _ensure_style(self, name: str, base_id: str, italic: bool = False) -> str:
        for sid, s in self.st_table.items():
            if s.name.lower() == name.lower():
                return sid
        st = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        base_name = self.st_table[base_id].name if base_id in self.st_table else "Normal"
        st.base_style = self.doc.styles[base_name]
        if italic:
            st.font.italic = True
        self.st_table = style_table(self.doc)
        return st.style_id

    def _template_bullet_num(self) -> str:
        """The numId the template's bullet paragraphs use, when its first
        level is a bullet; otherwise a new bullet definition."""
        sid = self.styles.bullet
        info = self.st_table.get(sid)
        cands = []
        for p in self.body.iter(qn("w:p")):
            if pstyle_id(p) != sid:
                continue
            np_ = p.find(qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:numId"))
            if np_ is not None:
                cands.append(np_.get(qn("w:val")))
        if info and info.num:
            cands.append(info.num[0])
        for c in cands:
            # reuse it only when it also defines the nested levels content can use
            if c and c != "0" and all(self.num.level_format(c, lv) == "bullet" for lv in range(3)):
                return c
        return self.num.bullet_definition(self.p.profile.fonts.get("body") or "Arial")

    def _direct_numpr(self, level: int):
        """A numPr written directly on a template heading of this level (so
        new headings number like the template's), or None."""
        sid = self.styles.h1 if level == 0 else self.styles.h2
        for p in self.body.iter(qn("w:p")):
            if pstyle_id(p) == sid and own_text(p).strip():
                np_ = p.find(qn("w:pPr") + "/" + qn("w:numPr"))
                nid = np_.find(qn("w:numId")) if np_ is not None else None
                if nid is not None and nid.get(qn("w:val")) != "0":
                    return copy.deepcopy(np_)
        return None

    # ------------------------------------------------------------ mapping
    def _is_heading(self, el, level: int | None = None) -> bool:
        if tag(el) != "p":
            return False
        sid = pstyle_id(el)
        if level is None:
            return sid in (self.styles.h1, self.styles.h2)
        return sid == (self.styles.h1 if level == 0 else self.styles.h2)

    def _map_template(self) -> list[THeading]:
        chapters: list[THeading] = []
        started = False
        for el in self.body.iterchildren():
            if self._is_heading(el, 0) and own_text(el).strip():
                chapters.append(THeading(el, 0, own_text(el).strip()))
                started = True
            elif started and self._is_heading(el, 1) and own_text(el).strip():
                chapters[-1].children.append(THeading(el, 1, own_text(el).strip()))
        if not chapters:
            raise DocxBuildError(f"template {self.template.name}: no Heading 1 paragraphs in style "
                                 f"'{self.st_table[self.styles.h1].name}'")
        return chapters

    def span_after(self, el, stop_levels: tuple[int, ...]) -> list:
        """Elements after a heading up to the next heading (of the given
        levels) with text, or the body's sectPr."""
        out = []
        nxt = el.getnext()
        while nxt is not None:
            if tag(nxt) == "sectPr":
                break
            if tag(nxt) == "p" and own_text(nxt).strip() and any(self._is_heading(nxt, lv) for lv in stop_levels):
                break
            out.append(nxt)
            nxt = nxt.getnext()
        return out

    def clear(self, els: list) -> list:
        """Remove guidance; returns the template tables found (detached)
        for cloning. Elements carrying a section break are kept, emptied."""
        protos = []
        for el in els:
            if el.getparent() is None:
                continue
            if has_sect(el):
                self.neutralise(el)
                continue
            if tag(el) == "tbl":
                protos.append(el)
            el.getparent().remove(el)
        return protos

    def neutralise(self, el) -> None:
        """Keep a paragraph that carries a section break, but make it an
        empty body paragraph (keeping a page break if it had one)."""
        if tag(el) != "p":
            # a table or content control holding a section break: keep only the break
            sect = copy.deepcopy(el.find(".//" + qn("w:sectPr")))
            p = OxmlElement("w:p")
            ppr = p.get_or_add_pPr()
            ppr.append(sect)
            el.addprevious(p)
            el.getparent().remove(el)
            el = p
        had_pb = any(b.get(qn("w:type")) == "page" for b in el.iter(qn("w:br")))
        ppr = el.find(qn("w:pPr"))
        for ch in list(el):
            if ch is not ppr:
                el.remove(ch)
        if ppr is not None:
            for ch in list(ppr):
                if tag(ch) != "sectPr":
                    ppr.remove(ch)
            ps = OxmlElement("w:pStyle")
            ps.set(qn("w:val"), self.styles.body)
            ppr.insert(0, ps)
        if had_pb:
            r = OxmlElement("w:r")
            r.append(make("w:br", w_type="page"))
            el.append(r)

    def set_heading(self, el, text: str, keep_run_format: bool = False) -> None:
        """Replace a paragraph's text. A heading's new run takes its
        formatting from the paragraph and its style alone, so a struck or
        highlighted template word never reaches the new title; cover text
        (``keep_run_format``) keeps its run formatting without guidance marks."""
        first_rpr = el.find(".//" + qn("w:r") + "/" + qn("w:rPr")) if keep_run_format else None
        rpr = clean_rpr(copy.deepcopy(first_rpr)) if first_rpr is not None else None
        for ch in list(el):
            if tag(ch) not in ("pPr",):
                el.remove(ch)
        r = _text_run(text)
        if rpr is not None:
            r.insert(0, rpr)
        el.append(r)

    def new_heading(self, level: int, text: str, like=None):
        sid = self.styles.h1 if level == 0 else self.styles.h2
        p = self.r.para(sid)
        p.add_run(text)
        src = None
        if like is not None:
            src = like.find(qn("w:pPr") + "/" + qn("w:numPr"))
        if src is None:
            src = self.direct_num[level]
        if src is not None:
            p._p.get_or_add_pPr().append(copy.deepcopy(src))
        return p._p

    # ------------------------------------------------------------ filling
    def _fill(self, tchapters: list[THeading]) -> None:
        chapters = self.p.chapters
        pairs = match(len(chapters), len(tchapters),
                      lambda i, j: chapter_score(chapters[i].title, [s.title for s in chapters[i].sections],
                                                 tchapters[j].title, [s.title for s in tchapters[j].children]),
                      0.4)
        by_ch = {i: j for i, j in pairs}
        used = {j for _, j in pairs}
        for j, tc in enumerate(tchapters):
            if j not in used:
                self.clear(self.span_after(tc.el, (0,)))
                if has_sect(tc.el):
                    self.neutralise(tc.el)
                else:
                    tc.el.getparent().remove(tc.el)
        prev_end = None  # the last element of the previous chapter
        for i, ch in enumerate(chapters):
            j = by_ch.get(i)
            if j is not None:
                head = tchapters[j].el
                self.set_heading(head, ch.title)
                self._fill_chapter(ch, head, tchapters[j].children)
            else:
                head = self.new_heading(0, ch.title)
                if prev_end is None:
                    first = next((tchapters[k].el for k in sorted(used)), None)
                    if first is not None:
                        first.addprevious(head)
                    else:
                        self.body.find(qn("w:sectPr")).addprevious(head)
                else:
                    prev_end.addnext(head)
                self._fill_chapter(ch, head, [])
            ch_span = self.span_after(head, (0,))
            prev_end = ch_span[-1] if ch_span else head

    def _fill_chapter(self, ch, head, tsections: list[THeading]) -> None:
        lead = self.span_after(head, (0, 1))
        protos = self.clear(lead)
        self.r.protos = protos
        a = head
        for el in self.r.blocks(ch.lead, head):
            a.addnext(el)
            a = el
        secs = ch.sections
        pairs = match(len(secs), len(tsections), lambda i, j: similarity(secs[i].title, tsections[j].title), 0.5)
        by_sec = {i: j for i, j in pairs}
        used = {j for _, j in pairs}
        for j, ts in enumerate(tsections):
            if j not in used:
                self.clear(self.span_after(ts.el, (0, 1)))
                if has_sect(ts.el):
                    self.neutralise(ts.el)
                else:
                    ts.el.getparent().remove(ts.el)
        last = None
        for i, s in enumerate(secs):
            j = by_sec.get(i)
            if j is not None:
                h = tsections[j].el
                self.set_heading(h, s.title)
                protos = self.clear(self.span_after(h, (0, 1)))
            else:
                h = self.new_heading(1, s.title, like=last)
                ref = last if last is not None else a
                if last is None:
                    # straight after the lead-in
                    a.addnext(h)
                else:
                    tail = self.span_after(ref, (0, 1))
                    (tail[-1] if tail else ref).addnext(h)
                protos = []
            self.r.protos = protos
            x = h
            for el in self.r.blocks(s.blocks, h):
                x.addnext(el)
                x = el
            last = h

    def _neutralise_empty_headings(self) -> None:
        for el in list(self.body.iterchildren()):
            if tag(el) == "p" and not own_text(el).strip() and \
                    paragraph_outline(el, self.st_table) in (0, 1):
                if el.find(".//" + qn("w:drawing")) is not None:
                    continue
                self.neutralise(el)

    # ------------------------------------------------------- front matter
    def _first_heading(self):
        for el in self.body.iterchildren():
            if self._is_heading(el, 0) and own_text(el).strip():
                return el
        return None

    def _front_matter(self) -> None:
        """Fill the document-control tables before the first chapter and
        remove bracketed guidance there."""
        first = self._first_heading()
        front = []
        for el in self.body.iterchildren():
            if el is first:
                break
            front.append(el)
        for el in front:
            if tag(el) == "p":
                t = own_text(el).strip()
                if re.fullmatch(r"\[[^\]]*\]", t) and el.find(".//" + qn("w:drawing")) is None:
                    if has_sect(el):
                        self.neutralise(el)
                    else:
                        el.getparent().remove(el)
            if tag(el) == "tbl":
                self._control_table(el)

    def _control_table(self, tbl) -> None:
        trs = tbl.findall(qn("w:tr"))
        if not trs:
            return
        header = [norm(ptext(tc)) for tc in trs[0].findall(qn("w:tc"))]
        hs = set(header)
        rows: list[list[str]] | None = None
        if any("version" in h for h in header) and any(
                k in h for h in header for k in ("date", "change", "rationale", "description", "summary")):
            rows = []
            for h in history_rows(self.p):
                row = []
                for col in header:
                    if "version" in col:
                        row.append(h["version"])
                    elif "date" in col:
                        row.append(h["date"])
                    elif any(k in col for k in ("author", "by", "name", "editor")):
                        row.append(h["author"])
                    else:
                        row.append(h["change"])
                rows.append(row)
        elif any(k in h for h in header for k in ("interest", "concern", "involvement", "business unit")):
            rows = self._stakeholder_rows(header)
        elif "name" in hs and any("role" in h or "position" in h or "title" == h for h in header):
            rows = []
            for label, who in (("Owner", self.p.config.get("owner")), ("Author", self.p.config.get("author"))):
                name, pos = _person(who)
                row, role_seen = [], False
                for col in header:
                    if col == "name":
                        row.append(name)
                    elif ("role" in col or col in ("title", "position")) and not role_seen:
                        role_seen = True
                        row.append(label)
                    else:
                        row.append(pos)
                rows.append(row)
        if rows is None:
            return
        proto = copy.deepcopy(trs[1] if len(trs) > 1 else trs[0])
        for tr in trs[1:]:
            tbl.remove(tr)
        for row in rows:
            tr = copy.deepcopy(proto)
            self.r._set_row(tr, row)
            tbl.append(tr)
        total = text_box(governing_sectpr(tbl, self.body))[0] // EMU_PER_TWIP
        if fit_table(self.doc, tbl, total, keep_if_fits=True, force=True) == "squeezed":
            self.r.squeezed.append(squeeze_warning("a document-control table", [ptext(tc) for tc in
                                                   trs[0].findall(qn("w:tc"))], False))
        finish_header(self.doc, tbl)

    def _stakeholder_rows(self, header: list[str]) -> list[list[str]]:
        reg = self.p.registers.get("stakeholder")
        if reg is None:
            return []
        t = reg.type
        fields = ["id", "title", "status"] + list(t.fields)
        out = []
        for rec in reg:
            row = []
            for col in header:
                syn = next((f for words, f in STAKEHOLDER_SYNONYMS if f in fields and any(w in col for w in words)),
                           None)
                if syn:
                    row.append(rec.display(syn))
                    continue
                best, score = None, 0.0
                for f in fields:
                    sc = max(similarity(col, f.replace("_", " ")), similarity(col, t.header(f)))
                    if col in ("name",) and f == "name":
                        sc = 1.0
                    if sc > score:
                        best, score = f, sc
                row.append(rec.display(best) if best and score >= 0.4 else "")
            out.append(row)
        return out

    def _fields(self) -> None:
        """Replace cover and front-matter text named in the profile's
        ``template.fields`` ({template text: replacement}); replacements
        may use placeholders such as ``{doc:title}``."""
        fields = self.p.profile.get("template.fields") or {}
        if not isinstance(fields, dict) or not fields:
            return
        values = {}
        for k, v in fields.items():
            if v is None:
                continue
            text, unresolved = self.p.resolve_text(str(v))
            for u in unresolved:
                self.warnings.append(f"template.fields '{k}': {u.token} {u.reason}")
            values[str(k)] = text
        for p in self.body.iter(qn("w:p")):
            if p.find(".//" + qn("w:p")) is not None:
                continue
            t = ptext(p)
            if t in values:
                self.set_heading(p, values[t], keep_run_format=True)

    def _clean_bookmarks(self) -> None:
        starts = {b.get(qn("w:id")) for b in self.body.iter(qn("w:bookmarkStart"))}
        ends = {b.get(qn("w:id")) for b in self.body.iter(qn("w:bookmarkEnd"))}
        for b in list(self.body.iter(qn("w:bookmarkStart"))):
            if b.get(qn("w:id")) not in ends:
                b.getparent().remove(b)
        for b in list(self.body.iter(qn("w:bookmarkEnd"))):
            if b.get(qn("w:id")) not in starts:
                b.getparent().remove(b)

    # ---------------------------------------------------------------- TOC
    def _toc(self) -> tuple[list[tuple[int, str | None, str, str]], set[str]]:
        """TOC entries for every outline-level 1 and 2 paragraph, as Word's
        own update would list them, with their rendered numbers."""
        expected = heading_entries(self.p)
        first_title = expected[0][2] if expected else None
        it = iter(expected)
        next_id = [_bookmark_ids(self.body)]
        entries, unnumbered, started = [], set(), False
        for p, lvl, label, _fmt, text in heading_paragraphs(self.doc):
            if not started and lvl == 1 and text == first_title:
                started = True
            if started:
                want = next(it, None)
                anchor = want[3] if want else f"x{len(entries)}"
            else:
                anchor = f"front{len(entries)}"
            bm = heading_bookmark(p, bookmark_name(anchor), next_id)
            if label is None and started:
                unnumbered.add(text)
            entries.append((lvl, label, text, bm))
        return entries, unnumbered

    def _warn_unnumbered_refs(self, unnumbered: set[str]) -> None:
        if not unnumbered:
            return
        for b in self.p.blocks():
            for raw in [b.raw] + list(getattr(b, "raw_items", [])):
                for m in re.finditer(r"\{sec:([^{}/]+)/([^{}]+)\}", raw or ""):
                    s = self.p.section(m.group(1).strip(), m.group(2).strip())
                    if s is not None and s.title in unnumbered:
                        self.warnings.append(f"{b.source}: {m.group(0)} gives {s.number}, but the template "
                                             f"leaves '{s.title}' unnumbered")


class TemplateRenderer(Renderer):
    """The renderer in template mode: tables clone a template table in the
    same section when the header matches."""

    protos: list = []
    allow_landscape = False

    def prototypes_available(self) -> bool:
        return bool(self.protos)

    def prototype_table(self, header: list[str]):
        want = [norm(C.plain(h)) for h in header]
        for proto in self.protos:
            trs = proto.findall(qn("w:tr"))
            if not trs:
                continue
            got = [norm(ptext(tc)) for tc in trs[0].findall(qn("w:tc"))]
            if got == want:
                return copy.deepcopy(proto)
        for proto in self.protos:
            trs = proto.findall(qn("w:tr"))
            if not trs:
                continue
            got = [ptext(tc) for tc in trs[0].findall(qn("w:tc"))]
            if len(got) == len(header) and len(trs) >= 1 and \
                    sum(similarity(a, b) for a, b in zip(got, header)) / len(header) >= 0.5:
                return copy.deepcopy(proto)
        return None


# Run properties that mark template guidance rather than format text.
GUIDANCE_MARKS = ("strike", "dstrike", "highlight", "shd", "vanish", "specVanish", "em")
# Run properties a cloned table cell keeps from the template's sample text: the face and size.
CELL_RPR_KEEP = ("rFonts", "sz", "szCs", "lang", "kern", "spacing", "w", "b", "bCs")


def clean_rpr(rpr):
    """Run properties without template guidance marks."""
    if rpr is None:
        return None
    for ch in list(rpr):
        if tag(ch) in GUIDANCE_MARKS:
            rpr.remove(ch)
    return rpr if len(rpr) else None


def keep_cell_rpr(rpr):
    """The run properties a new cell text keeps from a template sample cell:
    its face, size and weight. Colour, italics, strike-through and
    highlight there are the sample's guidance styling, not the table's."""
    if rpr is None:
        return None
    for ch in list(rpr):
        if tag(ch) not in CELL_RPR_KEEP:
            rpr.remove(ch)
    return rpr if len(rpr) else None


def clean_revisions_and_comments(doc) -> None:
    """Accept tracked insertions, drop deletions and property changes,
    remove comment anchors and the comments parts."""
    root = doc.element.body
    for ins in list(root.iter(qn("w:ins"))):
        par = ins.getparent()
        if par is None:
            continue
        if tag(par) == "rPr":
            par.remove(ins)
            continue
        idx = par.index(ins)
        for ch in list(ins):
            par.insert(idx, ch)
            idx += 1
        par.remove(ins)
    for name in ("w:del", "w:moveFrom", "w:rPrChange", "w:pPrChange", "w:sectPrChange", "w:tblPrChange",
                 "w:tcPrChange", "w:trPrChange", "w:tblGridChange", "w:numberingChange", "w:cellIns", "w:cellDel"):
        for el in list(root.iter(qn(name))):
            if el.getparent() is not None:
                el.getparent().remove(el)
    for el in list(root.iter(qn("w:moveTo"))):
        par = el.getparent()
        idx = par.index(el)
        for ch in list(el):
            par.insert(idx, ch)
            idx += 1
        par.remove(el)
    for name in ("w:moveFromRangeStart", "w:moveFromRangeEnd", "w:moveToRangeStart", "w:moveToRangeEnd",
                 "w:commentRangeStart", "w:commentRangeEnd"):
        for el in list(root.iter(qn(name))):
            el.getparent().remove(el)
    for ref in list(root.iter(qn("w:commentReference"))):
        r = ref.getparent()
        if tag(r) == "r":
            r.getparent().remove(r)
        else:
            r.remove(ref)
    for rid, rel in list(doc.part.rels.items()):
        if "comments" in rel.reltype or rel.reltype.endswith("/people"):
            del doc.part.rels[rid]
    # struck-through runs are template edits never made; highlights are review marks
    for r in list(root.iter(qn("w:r"))):
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            continue
        struck = [x for x in rpr if tag(x) in ("strike", "dstrike") and x.get(qn("w:val")) not in ("0", "false")]
        if struck and r.getparent() is not None and r.find(qn("w:drawing")) is None:
            r.getparent().remove(r)
    for h in list(root.iter(qn("w:highlight"))):
        h.getparent().remove(h)
    tr = doc.settings.element.find(qn("w:trackRevisions"))
    if tr is not None:
        doc.settings.element.remove(tr)


__all__ = ["DocxBuildError", "Structure", "align", "fill_gaps", "match", "build", "chapter_score", "detect_styles", "read_structure",
           "rebuild_toc", "similarity", "style_table"]
