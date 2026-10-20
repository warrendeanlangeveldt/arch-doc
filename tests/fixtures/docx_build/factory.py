"""Builders for the docx and extract tests: an organisation-style template
generated with python-docx, and a small project that uses every block kind.

The template has custom heading style names (outline levels inherited
through basedOn, as real templates do), heading numbering on the styles,
a cover with field text, document-control tables, a TOC field holding stale
entries, guidance paragraphs and sample rows, a landscape section, an empty
heading carrying a section break, unnumbered appendix sections, a chapter
the metamodel does not have, a tracked change and a comment.
"""
from __future__ import annotations

import json
import textwrap
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt

GUIDANCE = [
    "[Describe what this release delivers, in two or three sentences.]",
    "Guidance: list every decision the board is asked to make.",
    "[Write this section last.]",
    "Use the table below for each item inside the boundary.",
    "Explain how this design follows the blueprint.",
    "This chapter is retained for legacy reasons and should be removed.",
    "Record every architecture risk with its rating and owner.",
    "Define each term a reader may not know.",
    "Sample Org Risk Description",
    "Sample decision area",
]

TEMPLATE_CHAPTERS = [
    ("Executive Summary of the Architecture", [
        ("Delivery Scope", [GUIDANCE[0]], None),
        ("Decision Sought", [GUIDANCE[1]], (["Decision Area", "Decision Required"], [GUIDANCE[9], "Approve"])),
        ("Executive Summary", [GUIDANCE[2]], None),
    ]),
    ("Solution Context & Blueprint Alignment", [
        ("In Scope", [GUIDANCE[3]], (["Scope Item", "Description"], ["Example item", "Example description"])),
        ("Out of Scope", [], (["Out-of-scope Item", "Reason", "Future Treatment"], ["x", "y", "z"])),
        ("Parent Architecture Alignment", [GUIDANCE[4]], None),
    ]),
    ("Legacy Notes", [
        ("Old Material", [GUIDANCE[5]], None),
    ]),
    ("Risks, Issues and Dependencies", [
        ("Architecture Risks", [GUIDANCE[6]],
         (["Risk ID", "Risk", "Likelihood", "Impact", "Rating", "Owner", "Status"],
          ["R-1", GUIDANCE[8], "Low", "Low", "Low", "Someone", "Open"])),
        ("Operational Risks", [], None),
        ("Issues", [], None),
        ("Dependencies", [], None),
    ]),
    ("Appendix", [
        ("Appendix A - Glossary", [GUIDANCE[7]], (["Term", "Definition", "Domain"], ["t", "d", "x"])),
        ("Appendix B - Required Views", [], None),
    ]),
]


def _style(d, name, base, kind=WD_STYLE_TYPE.PARAGRAPH):
    st = d.styles.add_style(name, kind)
    st.base_style = d.styles[base]
    return st


def _heading_numbering(d, chapter_sid: str, section_sid: str) -> int:
    numbering = d.part.numbering_part.element
    levels = "".join(
        f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        + (f'<w:pStyle w:val="{[chapter_sid, section_sid][i]}"/>' if i < 2 else "")
        + f'<w:lvlText w:val="{".".join(f"%{k + 1}" for k in range(i + 1))}"/><w:lvlJc w:val="left"/></w:lvl>'
        for i in range(3))
    numbering.append(parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="77">'
                               f'<w:multiLevelType w:val="multilevel"/>{levels}</w:abstractNum>'))
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="77"><w:abstractNumId w:val="77"/></w:num>'))
    return 77


def _bullet_numbering(d) -> int:
    numbering = d.part.numbering_part.element
    an = parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="78"><w:multiLevelType w:val="hybridMultilevel"/>'
                   f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="-"/>'
                   f'<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl>'
                   f'</w:abstractNum>')
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(an)
    else:
        numbering.append(an)
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="78"><w:abstractNumId w:val="78"/></w:num>'))
    return 78


def _numpr(p, num_id: int, ilvl: int = 0) -> None:
    np_ = p._p.get_or_add_pPr().get_or_add_numPr()
    np_.get_or_add_ilvl().val = ilvl
    np_.get_or_add_numId().val = num_id


def _table(d, header, row, style="Org Table"):
    t = d.add_table(rows=2, cols=len(header))
    t.style = d.styles[style]
    for c, v in zip(t.rows[0].cells, header):
        c.text = v
    for c, v in zip(t.rows[1].cells, row):
        c.text = v
    # a distinctive grid width the fill must keep when it clones this table
    for gc in t._tbl.tblGrid.findall(qn("w:gridCol")):
        gc.set(qn("w:w"), "1333")
    return t


def _toc(d):
    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/>'
        f'<w:docPartUnique/></w:docPartObj></w:sdtPr><w:sdtContent>'
        f'<w:p><w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr><w:r><w:t>Contents</w:t></w:r></w:p>'
        f'<w:p><w:pPr><w:pStyle w:val="TOC1"/><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="9000"/></w:tabs>'
        f'</w:pPr><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
        f'<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        f'<w:hyperlink w:anchor="_Toc100"><w:r><w:t>1</w:t></w:r><w:r><w:tab/></w:r>'
        f'<w:r><w:t>Stale Chapter</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>3</w:t></w:r></w:hyperlink></w:p>'
        f'<w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
        f'</w:sdtContent></w:sdt>')
    d.element.body.find(qn("w:sectPr")).addprevious(sdt)


def make_template(path: Path) -> Path:
    """Write the test template to ``path``."""
    d = Document()
    chapter = _style(d, "Org Chapter", "Heading 1")
    section = _style(d, "Org Section", "Heading 2")
    body = _style(d, "Org Body", "Normal")
    body.font.size = Pt(10)
    bullet = _style(d, "Org Bullet", "List Paragraph")
    for name in ("toc 1", "toc 2"):
        st = _style(d, name, "Normal")
        st.element.set(qn("w:customStyle"), "0")
    tstyle = d.styles.add_style("Org Table", WD_STYLE_TYPE.TABLE)
    tstyle.base_style = d.styles["Table Grid"]
    num = _heading_numbering(d, chapter.style_id, section.style_id)
    for st, lvl in ((chapter, 0), (section, 1)):
        np_ = st.element.get_or_add_pPr().get_or_add_numPr()
        np_.get_or_add_ilvl().val = lvl
        np_.get_or_add_numId().val = num
    bnum = _bullet_numbering(d)

    # cover
    d.add_paragraph("Document Title Goes Here", style="Title")
    d.add_paragraph("V0.0", style="Org Body")
    d.add_paragraph("[Insert the sponsor's name.]", style="Org Body")
    d.add_page_break()
    # document control
    d.add_paragraph("Document Control", style="Org Body")
    _table(d, ["Role", "Name", "Role"], ["Author", "", ""])
    d.add_paragraph("Version History", style="Org Body")
    _table(d, ["Release Date", "Version", "Author", "Rationale of Change"], ["1 Jan 2020", "0.0", "A. N. Other",
                                                                            "Template issued"])
    _toc(d)
    d.add_page_break()

    for ci, (title, sections) in enumerate(TEMPLATE_CHAPTERS):
        d.add_paragraph(title, style="Org Chapter")
        if title.startswith("Executive"):
            p = d.add_paragraph("Lead-in guidance for the chapter.", style="Org Body")
            # a tracked insertion and deletion, and a comment
            ins = parse_xml(f'<w:ins {nsdecls("w")} w:id="901" w:author="Reviewer" w:date="2020-01-01T00:00:00Z">'
                            f'<w:r><w:t xml:space="preserve"> Inserted guidance.</w:t></w:r></w:ins>')
            dele = parse_xml(f'<w:del {nsdecls("w")} w:id="902" w:author="Reviewer" w:date="2020-01-01T00:00:00Z">'
                             f'<w:r><w:delText>Deleted guidance.</w:delText></w:r></w:del>')
            p._p.append(ins)
            p._p.append(dele)
            d.add_comment(p.runs, text="Reviewer comment on the template", author="Reviewer")
        for si, (stitle, paras, table) in enumerate(sections):
            h = d.add_paragraph(stitle, style="Org Section")
            if title == "Appendix":
                _numpr(h, 0)  # unnumbered appendix sections
            for t in paras:
                d.add_paragraph(t, style="Org Body")
            if stitle == "Delivery Scope":
                b = d.add_paragraph("Guidance bullet about the scope", style="Org Bullet")
                _numpr(b, bnum)
            if table:
                _table(d, *table)
            if title.startswith("Solution Context") and stitle == "In Scope":
                # a landscape section starts after this table; and an empty heading carries the break
                d.add_section(WD_SECTION.NEW_PAGE)
                s = d.sections[-1]
                s.orientation = WD_ORIENT.LANDSCAPE
                s.page_width, s.page_height = s.page_height, s.page_width
                empty = d.add_paragraph("", style="Org Section")
                empty._p.get_or_add_pPr().append(
                    parse_xml(f'<w:sectPr {nsdecls("w")}><w:pgSz w:w="16838" w:h="11906" w:orient="landscape"/>'
                              f'<w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" '
                              f'w:header="708" w:footer="708" w:gutter="0"/></w:sectPr>'))
                s.orientation = WD_ORIENT.PORTRAIT
                s.page_width, s.page_height = s.page_height, s.page_width
    d.save(str(path))
    return path


def make_png(path: Path, w: int, h: int, colour=(47, 111, 124)) -> Path:
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (w, h), "white")
    dr = ImageDraw.Draw(im)
    dr.rectangle([10, 10, w - 10, h - 10], outline=colour, width=6)
    dr.line([10, 10, w - 10, h - 10], fill=colour, width=4)
    im.save(path)
    return path


CONTENT = {
    "01-summary.md": """
        # summary
        This design delivers parcel **tracking**, with *care* and `code` and a [link](https://example.org/x).

        ## scope
        We deliver the service in {sec:context/in-scope}; see {fig:wide}.

        - First point with **bold**
          - Nested point
            - Deeper point
        - Second point

        1. Step one
        2. Step two
           1. Sub-step

        | Item | Description |
        |---|---|
        | Tracking API | The public *API* |
        | Scan feed | Carrier scans |

        ## decision-sought
        REGISTER decision-request

        ## overview
        FIGURE wide | A wide landscape figure
        FIGURE tall | A tall portrait figure
        """,
    "02-context.md": """
        # context
        ## in-scope
        | Scope Item | Description |
        |---|---|
        | Parcel events | Every scan, in order |
        ## out-of-scope
        Billing is out of scope.
        """,
    "19-risks.md": """
        # risks
        ## architecture-risks
        REGISTER risk where category=architecture
        VIEW ctx | Parcel context | solution-context
        ## operational-risks
        REGISTER risk where category=operational
        """,
    "22-appendix.md": """
        # appendix
        ## glossary
        REGISTER glossary-term
        ## object-guide
        CATALOGUE
        """,
}

REGISTERS = {
    "decision-request": """
        - {id: DRQ-01, title: Approve the design, area: Architecture, ask: approve, due: 2026-11-01, status: submitted, owner: Head of Parcels}
        """,
    "risk": """
        - {id: RSK-01, title: Carrier limits throttle scans, category: architecture, likelihood: medium, impact: high, rating: high, owner: Integration lead, status: open}
        - {id: RSK-02, title: Schema drift, category: architecture, likelihood: low, impact: medium, rating: medium, owner: Data lead, status: mitigated}
        - {id: RSK-03, title: On-call gap, category: operational, likelihood: medium, impact: medium, rating: medium, owner: Service owner, status: "TBC: Service owner"}
        """,
    "glossary-term": """
        - {id: GLO-01, title: Scan, definition: A carrier event for one parcel, domain: Operations, status: agreed}
        """,
    "stakeholder": """
        - {id: STK-01, title: Sponsor, name: Head of Parcels, team: Parcels, interest: Delivery on time, involvement: approver, status: engaged}
        """,
}


def write_project(root: Path, profile: dict, template: str | None = None, metamodel: str = "seed",
                  content: dict | None = None, figures: bool = True) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    cfg = {"title": "Parcel Tracking", "short": "PT", "kind": "solution", "version": "0.3",
           "date": "2026-10-20", "profile": "profile.yaml", "metamodel": metamodel, "template": template,
           "author": {"name": "Solution Architect", "role": "Architect", "team": "Parcels"},
           "owner": {"name": "Head of Parcels", "role": "Sponsor"},
           "history": [{"date": "2026-10-01", "version": "0.1", "author": "Solution Architect", "change": "First draft"},
                       {"date": "2026-10-20", "version": "0.3", "author": "Solution Architect", "change": "Review fixes"}]}
    (root / "arch-doc.yaml").write_text(yaml.safe_dump(cfg, sort_keys=False))
    (root / "profile.yaml").write_text(yaml.safe_dump(profile, sort_keys=False))
    (root / "content").mkdir(exist_ok=True)
    for fname, text in (CONTENT if content is None else content).items():
        (root / "content" / fname).write_text(textwrap.dedent(text).lstrip("\n"))
    (root / "registers").mkdir(exist_ok=True)
    for rtype, text in REGISTERS.items():
        (root / "registers" / f"{rtype}.yaml").write_text(textwrap.dedent(text).lstrip("\n"))
    ddir = root / "diagrams"
    ddir.mkdir(exist_ok=True)
    if figures:
        make_png(ddir / "wide.png", 1600, 800)
        make_png(ddir / "tall.png", 600, 2400)
        make_png(ddir / "ctx.png", 1200, 700)
        (ddir / "model.json").write_text(json.dumps({"objects": [
            {"id": "A1", "type": "Application", "name": "Tracking"},
            {"id": "A2", "type": "Application", "name": "Scan feed"},
            {"id": "D1", "type": "Data Object", "name": "Parcel event"}]}))
    return root
