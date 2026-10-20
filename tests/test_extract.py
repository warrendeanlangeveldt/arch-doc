"""Extraction (SPEC 4): a metamodel and a template fragment from a .docx."""
import sys
from pathlib import Path

import pytest
import yaml
from docx import Document

from conftest import northwind_profile, run_cli

sys.path.insert(0, str(Path(__file__).parent / "fixtures" / "docx_build"))
import factory  # noqa: E402

from archdoc import Project, load_metamodel, scaffold  # noqa: E402
from archdoc import docx_build as DB  # noqa: E402
from archdoc import extract as X  # noqa: E402


@pytest.fixture
def template(tmp_path):
    return factory.make_template(tmp_path / "org-template.docx")


def test_detects_styles_by_outline_level_not_name(template):
    styles = DB.detect_styles(Document(str(template)))
    assert styles["h1"] == "Org Chapter" and styles["h2"] == "Org Section"
    assert styles["body"] == "Org Body" and styles["bullet"] == "Org Bullet"
    assert styles["table"] == "Org Table"


def test_chapters_matched_to_seed(template):
    mm, frag, report = X.extract(template)
    keys = [c["key"] for c in mm["chapters"]]
    assert keys == ["summary", "context", "legacy-notes", "risks", "appendix"]
    by = {c["key"]: c for c in mm["chapters"]}
    assert by["summary"]["title"] == "Executive Summary of the Architecture"
    assert by["summary"]["picture"]["kind"] == "layers" and "unmatched" not in by["summary"]
    assert by["legacy-notes"]["unmatched"] is True
    assert by["legacy-notes"]["sections"][0]["unmatched"] is True
    assert "legacy reasons" in by["legacy-notes"]["sections"][0]["guidance"]
    # sections inherit the seed's keys, guidance and artefacts
    risks = {s["key"]: s for s in by["risks"]["sections"]}
    assert list(risks) == ["architecture-risks", "operational-risks", "issues", "dependencies"]
    assert risks["architecture-risks"]["artefacts"] == ["risk where category=architecture"]
    assert risks["architecture-risks"]["template_columns"][0][0] == "Risk ID"
    assert [s["key"] for s in by["appendix"]["sections"]][0] == "glossary"
    assert any("unmatched" in line for line in report)


def test_extracted_metamodel_loads(template, tmp_path):
    mm, _, _ = X.extract(template)
    path = tmp_path / "mm.yaml"
    path.write_text(X.dump(mm))
    loaded = load_metamodel(str(path))
    assert [c["key"] for c in loaded.chapters][0] == "summary"
    assert "risk" in loaded.artefact_types


def test_profile_fragment(template):
    _, frag, _ = X.extract(template)
    t = frag["template"]
    assert Path(t["path"]) == template.resolve()
    assert t["styles"]["h1"] == "Org Chapter" and t["styles"]["h2"] == "Org Section"
    assert t["fields"] == {"Document Title Goes Here": "{doc:title}", "V0.0": "V{doc:version}"}


def test_candidate_type_from_table_header():
    mm = load_metamodel("seed")
    t, cmap, score = X.candidate_type(["Risk ID", "Risk", "Cause", "Likelihood", "Impact", "Rating", "Owner"],
                                      "Delivery Risks", mm.artefact_types)
    assert t == "risk" and cmap["Likelihood"] == "likelihood" and cmap["Cause"] == "cause"
    assert X.candidate_type(["Colour", "Flavour"], "Ice cream", mm.artefact_types)[0] is None


def test_unmatched_section_proposes_register(tmp_path):
    d = Document()
    d.add_paragraph("Risks, Issues & Dependencies", style="Heading 1")
    d.add_paragraph("Commercial Exposure", style="Heading 2")
    t = d.add_table(rows=1, cols=5)
    for c, v in zip(t.rows[0].cells, ["Risk ID", "Risk", "Likelihood", "Impact", "Mitigation"]):
        c.text = v
    path = tmp_path / "t.docx"
    d.save(str(path))
    mm, _, _ = X.extract(path)
    sec = mm["chapters"][0]["sections"][0]
    assert sec["unmatched"] is True and sec["artefacts"] == ["risk"]
    assert sec["column_map"]["Mitigation"] == "mitigation"


def test_cli_writes_metamodel_and_fragment(template, tmp_path):
    mm, frag = tmp_path / "mm.yaml", tmp_path / "frag.yaml"
    r = run_cli("extract", template, "-o", mm, "--profile-out", frag)
    assert r.returncode == 0, r.stderr
    assert "legacy-notes (unmatched)" in r.stdout
    load_metamodel(str(mm))
    assert yaml.safe_load(frag.read_text())["template"]["styles"]["h1"] == "Org Chapter"
    r = run_cli("extract", template, "-o", tmp_path / "mm2.yaml")
    assert "template:" in r.stdout and "Org Section" in r.stdout


def test_cli_rejects_a_template_without_headings(tmp_path):
    d = Document()
    d.add_paragraph("No headings at all")
    path = tmp_path / "flat.docx"
    d.save(str(path))
    r = run_cli("extract", path, "-o", tmp_path / "mm.yaml")
    assert r.returncode == 2 and "no paragraphs at outline level 1" in r.stderr


def test_round_trip_extract_init_fill(template, tmp_path):
    """Generate a template, extract it, init a project on the extracted
    metamodel and fill the same template."""
    mm_path, frag_path = tmp_path / "mm.yaml", tmp_path / "frag.yaml"
    assert run_cli("extract", template, "-o", mm_path, "--profile-out", frag_path).returncode == 0
    prof = northwind_profile()
    prof.update(yaml.safe_load(frag_path.read_text()))
    prof_path = tmp_path / "profile.yaml"
    prof_path.write_text(yaml.safe_dump(prof))
    root = tmp_path / "rt"
    scaffold.init(root, kind="solution", profile=str(prof_path), metamodel=str(mm_path), title="Round Trip",
                  date="2026-10-20")
    p = Project.load(root)
    assert [c.key for c in p.chapters] == ["summary", "context", "legacy-notes", "risks", "appendix"]
    (root / "content" / "01-summary.md").write_text(
        "# summary\nWhat this design is for.\n## scope\nThe tracking service.\n## overview\n- One\n  - Two\n")
    (root / "content" / "03-legacy-notes.md").write_text("# legacy-notes\n## old-material\nKept on purpose.\n")
    p = Project.load(root)
    out = DB.build(p)[0]
    st = DB.read_structure(out)
    want = [(lvl, t) for lvl, _, t, _ in DB.heading_entries(p)]
    assert [(lvl, t) for lvl, _, t in st.headings] == want
    # the extracted titles are the template's own, so every template chapter is kept
    assert [t for lvl, _, t in st.headings if lvl == 1] == [c[0] for c in factory.TEMPLATE_CHAPTERS]
    assert st.headings[0] == (1, "1", "Executive Summary of the Architecture")
    assert st.sections == DB.read_structure(template).sections
    assert "Round Trip" in st.texts and "Kept on purpose." in st.texts
    for g in factory.GUIDANCE:
        assert g not in st.all_text()


def _doc(tmp_path, chapters):
    d = Document()
    for title, sections in chapters:
        d.add_paragraph(title, style="Heading 1")
        for stitle, paras, header in sections:
            d.add_paragraph(stitle, style="Heading 2")
            for t in paras:
                d.add_paragraph(t)
            if header:
                t = d.add_table(rows=1, cols=len(header))
                for c, v in zip(t.rows[0].cells, header):
                    c.text = v
    path = tmp_path / "t.docx"
    d.save(str(path))
    return path


def test_guidance_sections_go_to_views_or_are_set_aside(tmp_path):
    path = _doc(tmp_path, [
        ("Current State Architecture", [("Viewpoint Guidance", ["Show the views."], ["Required View", "Type"]),
                                        ("Current Application Landscape", [], None)]),
        ("Runbooks & Operational Procedures", [("Repository Guidance", ["Link the runbooks."], None),
                                               ("Required Runbooks", [], None)]),
    ])
    mm, _, report = X.extract(path)
    cur, run = mm["chapters"]
    assert [s["key"] for s in cur["sections"]] == ["views", "application-landscape"]
    assert cur["sections"][0]["template_columns"] == [["Required View", "Type"]]
    # runbooks has no views section in the seed: the guidance is kept beside the chapter, not as content
    assert [s["key"] for s in run["sections"]] == ["required"]
    assert run["template_guidance"] == [{"title": "Repository Guidance", "guidance": "Link the runbooks."}]


def test_unmatched_section_keeps_the_templates_own_guidance(tmp_path):
    path = _doc(tmp_path, [("Requirements & Design Drivers", [
        ("Business Requirements", [], None),
        ("Vendor Obligations", ["[List what each vendor must deliver.]"], None),
        ("Design Decisions", [], ["ID", "Decision", "Owner  "]),
    ])])
    mm, _, _ = X.extract(path)
    secs = {s["title"]: s for s in mm["chapters"][0]["sections"]}
    assert secs["Vendor Obligations"]["guidance"] == "List what each vendor must deliver."
    assert secs["Design Decisions"]["guidance"] == "The template's table: ID, Decision, Owner."
    assert secs["Design Decisions"]["template_columns"] == [["ID", "Decision", "Owner  "]]  # as written


def test_matched_sections_map_template_headers_to_fields(template):
    mm, _, _ = X.extract(template)
    risks = {s["key"]: s for c in mm["chapters"] if c["key"] == "risks" for s in c["sections"]}
    cmap = risks["architecture-risks"]["column_map"]
    assert cmap["Risk ID"] == "id" and cmap["Risk"] == "title" and cmap["Likelihood"] == "likelihood"
    load_metamodel  # the extracted metamodel, column_map included, validates (checked in extract)


def test_column_map_conventions():
    mm = load_metamodel("seed")
    t = mm.artefact_types["business-requirement"]
    assert X.column_map_for(["ID", "Business Requirement", "MoSCoW", "Design Response"], [t]) == \
        {"ID": "id", "Business Requirement": "title", "MoSCoW": "priority", "Design Response": "response"}


def test_header_matching_reads_slashes_containment_and_aliases():
    mm = load_metamodel("seed")
    T = mm.artefact_types
    st = X.column_map_for(["Stakeholder", "Role", "Business Unit/Team", "Interest/Concern"], [T["stakeholder"]],
                          {"name": ["stakeholder"]})
    assert st == {"Stakeholder": "name", "Role": "title", "Business Unit/Team": "team",
                  "Interest/Concern": "interest"}
    risk = X.column_map_for(["Risk ID", "Risk", "Risk Owner", "Due Date"], [T["risk"]])
    assert risk == {"Risk ID": "id", "Risk": "title", "Risk Owner": "owner"}
    glo = X.column_map_for(["Term / Acronym", "Definition", "Business Context / Domain", "Source / Owner"],
                           [T["glossary-term"]])
    assert glo == {"Term / Acronym": "title", "Definition": "definition", "Business Context / Domain": "domain",
                   "Source / Owner": "source"}
    # the first text column names the record when no header claims the title
    ready = X.column_map_for(["Readiness Area", "Status", "Evidence", "Owner"], [T["readiness-item"]])
    assert ready["Readiness Area"] == "title"
    flow = X.column_map_for(["Flow ID", "Data Entity", "Source Application", "Target Application"], [T["data-flow"]])
    assert flow == {"Flow ID": "id", "Data Entity": "data", "Source Application": "source",
                    "Target Application": "target"}


def test_extract_reads_field_aliases_from_the_metamodel(tmp_path):
    import copy
    seed = copy.deepcopy(load_metamodel("seed").raw)
    seed["artefact_types"]["cutover-step"]["aliases"] = {"order": ["step"]}
    sp = tmp_path / "seed-with-aliases.yaml"
    sp.write_text(yaml.safe_dump(seed, sort_keys=False))
    path = _doc(tmp_path, [("Cutover, Release, Hypercare & Handover",
                            [("Cutover Plan Summary", [], ["Step", "Activity", "Owner", "Start", "Finish"])])])
    mm, _, _ = X.extract(path, str(sp))
    sec = mm["chapters"][0]["sections"][0]
    assert sec["column_map"]["Step"] == "order" and sec["column_map"]["Activity"] == "title"


def test_stale_value_labels_name_extract_and_the_field():
    import copy
    from archdoc import metamodel as M
    d = copy.deepcopy(load_metamodel("seed").raw)
    d["extracted_from"] = "org.docx"
    d["artefact_types"]["integration"]["fields"]["pattern"] = "text"
    with pytest.raises(M.MetamodelError) as e:
        M.from_dict(d)
    msg = " ".join(e.value.errors)
    assert "archdoc extract copied the seed's value_labels" in msg and "integration.pattern" in msg


def test_a_label_match_beats_containment_and_the_type_label_names_the_title():
    T = load_metamodel("seed").artefact_types
    dev = X.column_map_for(["Deviation ID", "Blueprint Requirement / Decision", "Proposed Deviation", "Reason",
                            "Impact", "Approval Required"], [T["exception"]])
    assert dev["Blueprint Requirement / Decision"] == "standard" and dev["Proposed Deviation"] == "title"
    assert X.column_map_for(["Exception ID", "Required View", "Impact", "Compensating Action"],
                            [T["exception"]])["Compensating Action"] == "control"


def test_trailing_punctuation_is_ignored_and_headers_are_kept_as_written():
    import copy
    from archdoc.registers import ArtefactType
    raw = copy.deepcopy(load_metamodel("seed").raw["artefact_types"]["risk"])
    raw["fields"]["acceptance"] = "text"
    t, errs = ArtefactType.from_dict("risk", raw)
    assert not errs
    cmap = X.column_map_for(["Risk ID", "Risk", "Acceptance Required?", "Target "], [t])
    assert cmap["Acceptance Required?"] == "acceptance"


def test_extract_writes_value_labels_only_for_enum_values(tmp_path):
    import copy
    seed = copy.deepcopy(load_metamodel("seed").raw)
    seed["artefact_types"]["integration"]["fields"]["pattern"] = "text"  # a target where pattern is free text
    seed["artefact_types"]["integration"]["value_labels"] = {"api": "API", "etl": "ETL"}
    sp = tmp_path / "target.yaml"
    sp.write_text(yaml.safe_dump(seed, sort_keys=False))
    path = _doc(tmp_path, [("Integration & API Architecture", [("Integration Overview", [], None)])])
    mm, _, _ = X.extract(path, str(sp))
    assert "value_labels" not in mm["artefact_types"]["integration"]
    load_metamodel  # the result validates (extract checks it)


def test_one_broken_type_is_reported_once():
    import copy
    from archdoc import metamodel as M
    d = copy.deepcopy(load_metamodel("seed").raw)
    d["artefact_types"]["integration"]["fields"]["pattern"] = "text"
    errs = M.validate(d)
    assert errs and all("artefact_types.integration.value_labels" in e for e in errs)
