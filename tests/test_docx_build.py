"""The Word build (SPEC 6.1): neutral document and template fill."""
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
import yaml
from docx import Document
from docx.oxml.ns import qn

from conftest import EXAMPLES, northwind_profile, run_cli

sys.path.insert(0, str(Path(__file__).parent / "fixtures" / "docx_build"))
import factory  # noqa: E402

from archdoc import Project  # noqa: E402
from archdoc import docx_build as DB  # noqa: E402


def expected_headings(p):
    return [(lvl, num, title) for lvl, num, title, _ in DB.heading_entries(p)]


@pytest.fixture
def neutral(tmp_path):
    root = factory.write_project(tmp_path / "doc", northwind_profile())
    p = Project.load(root)
    assert p.issues == []
    out = DB.build(p)
    assert out == [root / "out" / "Parcel Tracking v0.3.docx"]
    return p, out[0]


@pytest.fixture
def template(tmp_path):
    return factory.make_template(tmp_path / "org-template.docx")


def fill(tmp_path, template, profile_extra=None, **kw):
    prof = northwind_profile()
    prof.update(profile_extra or {})
    root = factory.write_project(tmp_path / "filled", prof, template=str(template), **kw)
    p = Project.load(root)
    return p, DB.build(p)[0]


# ------------------------------------------------------------------ neutral

def test_neutral_numbering_and_toc(neutral):
    p, out = neutral
    st = DB.read_structure(out)
    assert st.headings == expected_headings(p)
    assert st.toc == expected_headings(p)
    assert st.update_fields and st.toc_dirty
    assert st.headings[0] == (1, "1", "Executive Architecture Summary")
    assert (2, "19.2", "Operational Risks") in st.headings


def test_neutral_cover_and_document_control(neutral):
    p, out = neutral
    st = DB.read_structure(out)
    assert "Parcel Tracking" in st.texts and "Solution design" in st.texts
    assert ["Owner", "Head of Parcels", "Sponsor"] in st.tables[0]
    assert st.tables[1][1:] == [["1 October 2026", "0.1", "Solution Architect", "First draft"],
                                ["20 October 2026", "0.3", "Solution Architect", "Review fixes"]]
    d = Document(str(out))
    assert d.core_properties.title == "Parcel Tracking"
    assert "PT v0.3" in d.sections[0].footer.paragraphs[0].text


def test_neutral_renders_every_block_kind(neutral):
    p, out = neutral
    st = DB.read_structure(out)
    text = st.all_text()
    # paragraphs, placeholders resolved
    assert "We deliver the service in 2.1; see Figure 1." in st.texts
    # nested bullets and ordered lists with their own numbering
    assert (0, "bullet", "First point with bold") in st.list_items
    assert (1, "bullet", "Nested point") in st.list_items
    assert (2, "bullet", "Deeper point") in st.list_items
    assert (0, "decimal", "Step one") in st.list_items and (1, "lowerLetter", "Sub-step") in st.list_items
    # plain table, register tables, catalogue
    assert [["Item", "Description"], ["Tracking API", "The public API"], ["Scan feed", "Carrier scans"]] in st.tables
    risks = next(t for t in st.tables if t[0][:2] == ["ID", "Risk"])
    assert [r[0] for r in risks[1:]] == ["RSK-01", "RSK-02"]
    assert risks[1][-1] == "Open" and risks[2][-1] == "Mitigated"
    op = next(t for t in st.tables if t[1:] and t[1][0] == "RSK-03")
    # the status cell is whatever the core displays for the record (a TBC stays a TBC)
    core = next(b for b in p.blocks() if getattr(b, "type", None) == "risk" and b.records[0].id == "RSK-03")
    assert op[1] == core.compact_rows[0] and op[1][-1].startswith("TBC")
    assert [["Object type", "Objects", "Examples"], ["Application", "2", "Tracking; Scan feed"],
            ["Data Object", "1", "Parcel event"]] in st.tables
    # figures and views with core numbering
    assert st.captions == ["Figure 1. A wide landscape figure", "Figure 2. A tall portrait figure",
                           "View 1. Parcel context (ArchiMate Layered)"]
    assert len(st.pictures) == 3
    assert "{" not in text


def test_inline_styles_and_status_chips(neutral):
    p, out = neutral
    d = Document(str(out))
    runs = [r for para in d.paragraphs for r in para.runs]
    assert any(r.text == "tracking" and r.bold for r in runs)
    assert any(r.text == "care" and r.italic for r in runs)
    assert any(r.text == "code" and r.font.name == p.profile.fonts["mono"] for r in runs)
    body = d.element.body
    links = [h for h in body.iter(qn("w:hyperlink")) if h.get(qn("r:id"))]
    assert any(d.part.rels[h.get(qn("r:id"))].target_ref == "https://example.org/x" for h in links)
    # the status cell is a word, coloured by its chip
    for t in d.tables:
        if t.rows[0].cells[0].text == "ID" and t.rows[0].cells[1].text == "Risk":
            cell = t.rows[1].cells[-1].paragraphs[0].runs[0]
            assert cell.text == "Open" and cell.bold
            assert str(cell.font.color.rgb) == p.profile.palette["not_started"].lstrip("#").upper()
            break
    else:
        pytest.fail("risk table not found")


def test_neutral_heading_numbering_is_on_the_styles(neutral):
    p, out = neutral
    d = Document(str(out))
    for name, lvl in (("Heading 1", "0"), ("Heading 2", "1")):
        np_ = d.styles[name].element.pPr.find(qn("w:numPr"))
        assert np_.find(qn("w:ilvl")).get(qn("w:val")) == lvl
    assert d.styles["Heading 1"].font.name == p.profile.fonts["heading"]
    assert d.styles["Normal"].font.name == p.profile.fonts["body"]
    assert any(s.name == "House Table" for s in d.styles)


def test_no_guidance_left_in_neutral(neutral):
    p, out = neutral
    text = DB.read_structure(out).all_text()
    for ch in p.chapters:
        for s in ch.sections:
            assert s.guidance[:60] not in text


def test_figures_fit_text_width_and_page_height(neutral):
    p, out = neutral
    d = Document(str(out))
    sec = d.sections[0]
    text_w = sec.page_width - sec.left_margin - sec.right_margin
    text_h = sec.page_height - sec.top_margin - sec.bottom_margin
    pics = DB.read_structure(out).pictures
    wide, tall, view = pics
    assert abs(wide[0] - text_w) < 2000 and abs(view[0] - text_w) < 2000
    assert wide[1] == pytest.approx(wide[0] / 2, rel=0.01)
    assert tall[1] <= text_h - DB.CAPTION_ALLOWANCE and tall[0] < text_w
    assert tall[1] / tall[0] == pytest.approx(4, rel=0.01)


def test_missing_figure_fails_with_its_name(tmp_path):
    content = {"01-summary.md": "# summary\n## scope\nFIGURE nothing-here | Not drawn yet\n"}
    root = factory.write_project(tmp_path / "doc", northwind_profile(), content=content)
    with pytest.raises(DB.DocxBuildError, match="nothing-here"):
        DB.build(Project.load(root))
    assert not (root / "out" / "Parcel Tracking v0.3.docx").exists()


def test_mermaid_without_node_names_the_tool(tmp_path, monkeypatch):
    content = {"01-summary.md": "# summary\n## scope\nFIGURE flow | The flow\n"}
    root = factory.write_project(tmp_path / "doc", northwind_profile(), content=content)
    (root / "diagrams" / "flow.mmd").write_text("flowchart LR\n  a --> b\n")
    from archdoc import diagrams
    monkeypatch.setattr(diagrams, "find_npx", lambda: None)
    with pytest.raises(DB.DocxBuildError, match=r"'flow' cannot be embedded.*npx.*mermaid-cli"):
        DB.build(Project.load(root))


def test_ead_source_without_arch_diagram_names_it(tmp_path, monkeypatch):
    content = {"01-summary.md": "# summary\n## scope\nFIGURE ctx2 | Context\n"}
    root = factory.write_project(tmp_path / "doc", northwind_profile(), content=content)
    (root / "diagrams" / "ctx2.ead").write_text("view x\n")
    from archdoc import diagrams
    monkeypatch.setattr(diagrams, "find_arch_diagram", lambda *a, **k: None)
    with pytest.raises(DB.DocxBuildError, match="'ctx2' cannot be embedded.*arch-diagram"):
        DB.build(Project.load(root))


def test_empty_register_says_none_recorded(tmp_path):
    content = {"19-risks.md": "# risks\n## issues\nREGISTER issue\n"}
    root = factory.write_project(tmp_path / "doc", northwind_profile(), content=content)
    out = DB.build(Project.load(root))[0]
    t = next(t for t in DB.read_structure(out).tables if t[0][0] == "ID" and "None recorded." in t[1])
    assert len(t) == 2


def test_cli_build_docx(tmp_path):
    root = factory.write_project(tmp_path / "doc", northwind_profile())
    r = run_cli("-C", root, "build", "docx", "--force")
    assert r.returncode == 0, r.stderr
    assert (root / "out" / "Parcel Tracking v0.3.docx").exists()


# ------------------------------------------------------------- template fill

FIELDS = {"template": {"fields": {"Document Title Goes Here": "{doc:title}", "V0.0": "V{doc:version}"}}}


def test_template_keeps_sections_and_cleans_revisions(tmp_path, template):
    before = DB.read_structure(template)
    assert before.revisions and before.comments
    p, out = fill(tmp_path, template, FIELDS)
    st = DB.read_structure(out)
    assert st.sections == before.sections == 3
    assert st.revisions == 0 and st.comments == 0
    with zipfile.ZipFile(out) as z:
        assert not [n for n in z.namelist() if "comments" in n]
    assert "Inserted guidance" not in st.all_text() and "Deleted guidance" not in st.all_text()


def test_template_numbering_toc_and_unnumbered_appendix(tmp_path, template):
    p, out = fill(tmp_path, template, FIELDS)
    st = DB.read_structure(out)
    want = expected_headings(p)
    appendix = [t for lvl, n, t in want if lvl == 2 and n.startswith(str(p.chapter("appendix").number) + ".")]
    got = st.headings
    assert [(lvl, t) for lvl, _, t in got] == [(lvl, t) for lvl, _, t in want]
    for (lvl, n, t), (_, wn, _) in zip(got, want):
        assert n == wn or (n is None and t in appendix), (t, n, wn)
    assert all(n is None for lvl, n, t in got if t in appendix)
    assert st.toc == got
    assert st.update_fields and st.toc_dirty
    assert "Stale Chapter" not in [t for _, _, t in st.toc]


def test_template_uses_its_own_styles(tmp_path, template):
    p, out = fill(tmp_path, template, FIELDS)
    d = Document(str(out))
    styles = {para.style.name for para in d.paragraphs if para.text == p.chapters[0].title}
    assert styles == {"Org Chapter"}
    assert {para.style.name for para in d.paragraphs if para.text == "Delivery Scope"} == {"Org Section"}
    assert any(para.style.name == "Org Body" and para.text.startswith("This design delivers") for para in d.paragraphs)


def test_template_leaves_no_guidance(tmp_path, template):
    p, out = fill(tmp_path, template, FIELDS)
    st = DB.read_structure(out)
    text = st.all_text() + "\n".join(t for _, _, t in st.headings)
    for g in factory.GUIDANCE:
        assert g not in text, g
    for gone in ("Legacy Notes", "Old Material", "Guidance bullet about the scope", "Example item",
                 "[Insert the sponsor", "Lead-in guidance", "A. N. Other", "Template issued"):
        assert gone not in text, gone


def test_template_fills_document_control_and_fields(tmp_path, template):
    p, out = fill(tmp_path, template, FIELDS)
    st = DB.read_structure(out)
    assert st.tables[0] == [["Role", "Name", "Role"], ["Owner", "Head of Parcels", "Sponsor"],
                            # the role "Architect" is already in the name, so the position is the team
                            ["Author", "Solution Architect", "Parcels"]]
    assert st.tables[1][1:] == [["1 October 2026", "0.1", "Solution Architect", "First draft"],
                                ["20 October 2026", "0.3", "Solution Architect", "Review fixes"]]
    assert "Parcel Tracking" in st.texts and "V0.3" in st.texts
    assert "Document Title Goes Here" not in st.texts and "V0.0" not in st.texts


def test_template_clones_matching_tables(tmp_path, template):
    p, out = fill(tmp_path, template, FIELDS)
    d = Document(str(out))
    for t in d.tables:
        hdr = [c.text for c in t.rows[0].cells]
        if hdr == ["Scope Item", "Description"]:
            # the template's own table, cloned; its 1333-twip columns cannot hold "Description" at the
            # table font, so the widths are refitted to the content
            assert t.style.name == "Org Table"
            widths = [int(gc.get(qn("w:w"))) for gc in t._tbl.tblGrid.findall(qn("w:gridCol"))]
            mins, _ = DB.column_needs(d, t._tbl, DB.table_font_pt(d, t._tbl))
            assert all(w >= m for w, m in zip(widths, mins)) and widths != [1333, 1333]
            assert [c.text for c in t.rows[1].cells] == ["Parcel events", "Every scan, in order"]
            break
    else:
        pytest.fail("in-scope table not found")
    # a table with no template counterpart takes the template's table style
    item = next(t for t in d.tables if [c.text for c in t.rows[0].cells] == ["Item", "Description"])
    assert item.style.name == "Org Table"


def test_template_renders_every_block_kind(tmp_path, template):
    p, out = fill(tmp_path, template, FIELDS)
    st = DB.read_structure(out)
    assert st.captions == ["Figure 1. A wide landscape figure", "Figure 2. A tall portrait figure",
                           "View 1. Parcel context (ArchiMate Layered)"]
    assert (1, "bullet", "Nested point") in st.list_items and (0, "decimal", "Step one") in st.list_items
    assert [["Object type", "Objects", "Examples"], ["Application", "2", "Tracking; Scan feed"],
            ["Data Object", "1", "Parcel event"]] in st.tables
    assert any(r[:1] == ["RSK-01"] for t in st.tables for r in t)


def test_page_geometry_follows_the_governing_section(template):
    """A figure placed inside a landscape section is sized to that
    section's text width, not the document's."""
    d = Document(str(template))
    body = d.element.body
    breaks = [p for p in body.iter(qn("w:p")) if p.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None]
    landscape = next(p for p in breaks
                     if p.find(".//" + qn("w:pgSz")).get(qn("w:orient")) == "landscape")
    inside = landscape.getprevious()
    while inside is not None and inside in breaks:
        inside = inside.getprevious()
    portrait_w = DB.text_box(DB.governing_sectpr(breaks[0].getprevious(), body))[0]
    landscape_w = DB.text_box(DB.governing_sectpr(landscape, body))[0]
    assert landscape_w == (16838 - 2880) * DB.EMU_PER_TWIP
    assert landscape_w > portrait_w
    assert DB.fit_figure(1600, 800, DB.text_box(DB.governing_sectpr(landscape, body))) == landscape_w


def test_template_without_headings_fails(tmp_path):
    d = Document()
    d.add_paragraph("Just text")
    path = tmp_path / "flat.docx"
    d.save(str(path))
    with pytest.raises(DB.DocxBuildError, match="Heading 1"):
        fill(tmp_path, path)


def test_template_from_profile_path(tmp_path, template):
    prof = northwind_profile()
    prof["template"] = {"path": str(template), **FIELDS["template"]}
    root = factory.write_project(tmp_path / "viaprofile", prof)
    out = DB.build(Project.load(root))[0]
    assert "Org Chapter" in DB.read_structure(out).styles.values()


# ------------------------------------------------------------------ helpers

def test_similarity_and_align():
    assert DB.similarity("Risks, Issues & Dependencies", "Risks, Issues and Dependencies") == 1.0
    assert DB.similarity("Appendix A - Glossary", "Glossary") == 1.0
    assert DB.similarity("Security & Controls", "Data Migration") < 0.4
    pairs = DB.match(3, 3, lambda i, j: [[1, 0, 0], [0, 0.2, 0], [0, 0, 1]][i][j], 0.5)
    assert pairs == [(0, 0), (1, 1), (2, 2)]  # the middle one is filled by position


def test_examples_build_in_both_modes(tmp_path, examples_copy, template):
    for name in ("northwind-platform", "northwind-parcels"):
        root = examples_copy / name
        p = Project.load(root)
        neutral = DB.build(p, tmp_path / name / "neutral")[0]
        assert DB.read_structure(neutral).headings == expected_headings(p)
        cfg = yaml.safe_load((root / "arch-doc.yaml").read_text())
        cfg["template"] = str(template)
        (root / "arch-doc.yaml").write_text(yaml.safe_dump(cfg, sort_keys=False))
        p = Project.load(root)
        filled = DB.build(p, tmp_path / name / "template")[0]
        st = DB.read_structure(filled)
        # the template's three sections, plus a pair per table set in its own landscape section
        assert st.sections >= 3 and (st.sections - 3) % 2 == 0
        assert [t for _, _, t in st.headings] == [t for _, _, t in expected_headings(p)]


# ------------------------------------------------------------ table fitting

def _table(header, rows, grid=None):
    d = Document()
    t = d.add_table(rows=1 + len(rows), cols=len(header))
    for r, vals in zip(t.rows, [header] + rows):
        for c, v in zip(r.cells, vals):
            c.text = v
    if grid:
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), grid):
            gc.set(qn("w:w"), str(w))
    return d, t


def widths_of(t):
    return [int(gc.get(qn("w:w"))) for gc in t._tbl.tblGrid.findall(qn("w:gridCol"))]


def test_fit_gives_every_column_its_longest_token():
    d, t = _table(["ID", "Category", "Requirement", "Design response", "Status"],
                  [["NFR-01", "availability", "Tracking API availability",
                    "Three replicas or more across zones; read replicas of the store in the paired region "
                    "so a regional outage does not stop tracking", "Implemented"]])
    mins, nats = DB.column_needs(d, t._tbl, DB.table_font_pt(d, t._tbl))
    total = int(sum(mins)) + 3000
    assert sum(nats) > total  # the long column must wrap
    assert DB.fit_table(d, t._tbl, total) == "fit"
    w = widths_of(t)
    assert sum(w) == total and all(a >= m for a, m in zip(w, mins))
    assert w[3] == max(w)  # the long-text column takes the rest
    tblpr = t._tbl.tblPr
    assert tblpr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
    assert tblpr.find(qn("w:tblW")).get(qn("w:w")) == str(total)
    for row in t.rows:
        assert [int(c._tc.tcPr.find(qn("w:tcW")).get(qn("w:w"))) for c in row.cells] == w


def test_fit_steps_the_font_down_then_asks_for_landscape():
    header = [f"Column{i}" for i in range(7)]
    d, t = _table(header, [["Implementation"] * 7])
    base = DB.table_font_pt(d, t._tbl)
    at_base = sum(DB.column_needs(d, t._tbl, base)[0])
    at_min = sum(DB.column_needs(d, t._tbl, DB.MIN_TABLE_PT)[0])
    assert DB.fit_table(d, t._tbl, int(at_min) - 10) == "landscape"
    assert DB.fit_table(d, t._tbl, int((at_base + at_min) / 2)) == "fit"
    sizes = {r.font.size.pt for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs}
    assert len(sizes) == 1 and DB.MIN_TABLE_PT <= sizes.pop() < DB.table_font_pt(d, t._tbl)


def test_fit_keeps_template_widths_that_hold_the_content():
    d, t = _table(["ID", "Title"], [["A-1", "Short"]], grid=[2000, 7000])
    assert DB.fit_table(d, t._tbl, 9000, keep_if_fits=True) == "kept"
    assert widths_of(t) == [2000, 7000]
    d, t = _table(["ID", "Title"], [["Implementation-reference", "Short"]], grid=[500, 8500])
    assert DB.fit_table(d, t._tbl, 9000, keep_if_fits=True) == "fit"
    assert widths_of(t)[0] > 500


def test_wide_register_gets_a_landscape_section(tmp_path):
    content = {"19-risks.md": "# risks\n## architecture-risks\n"
               "REGISTER risk cols id,title,category,cause,consequence,likelihood,impact,rating,mitigation,"
               "residual,owner,status\nAfter the table.\n"}
    regs = {"risk": "- {id: RSK-01, title: Throttling, category: architecture, cause: Carrier-imposed-ratelimits, "
            "consequence: Delayed-notifications, likelihood: medium, impact: high, rating: high, "
            "mitigation: Queueing, residual: medium, owner: Integration-lead, status: mitigating}\n"}
    root = factory.write_project(tmp_path / "doc", northwind_profile(), content=content)
    (root / "registers" / "risk.yaml").write_text(regs["risk"])
    p = Project.load(root)
    out = DB.build(p)[0]
    d = Document(str(out))
    orients = [s.orientation for s in d.sections]
    from docx.enum.section import WD_ORIENT
    assert orients.count(WD_ORIENT.LANDSCAPE) == 1
    land = next(s for s in d.sections if s.orientation == WD_ORIENT.LANDSCAPE)
    assert land.page_width > land.page_height
    # every section after the cover keeps the footer
    assert all(s.footer.paragraphs[0].text.startswith("PT v0.3") for s in d.sections[1:])


def test_every_table_holds_its_longest_token(neutral, tmp_path, template):
    p, out = neutral
    assert DB.read_structure(out).narrow_tables == []
    _, filled = fill(tmp_path, template, FIELDS)
    assert DB.read_structure(filled).narrow_tables == []


def test_identifier_hyphens_do_not_break(neutral):
    p, out = neutral
    d = Document(str(out))
    body = d.element.body
    assert body.find(".//" + qn("w:noBreakHyphen")) is not None
    assert DB.identifier_tokens("RSK-04 and pts-prd-tracking-db, but zone-redundant") == \
        ["RSK-04", "pts-prd-tracking-db,"]
    st = DB.read_structure(out)
    assert any(r[:1] == ["RSK-01"] for t in st.tables for r in t)  # read back with its hyphen


@pytest.mark.skipif(not (shutil.which("soffice") and shutil.which("pdftotext")),
                    reason="LibreOffice (soffice) and pdftotext render the pages")
def test_rendered_tables_never_break_inside_a_word(neutral, tmp_path, template):
    """Render through LibreOffice and check every table word appears whole
    in the PDF text. Identifiers must stay whole; a compound word may break
    after its hyphen, as prose does."""
    _, filled = fill(tmp_path, template, FIELDS)
    for out in (neutral[1], filled):
        d = Document(str(out))
        words = set()
        for tbl in d.element.body.iter(qn("w:tbl")):
            for para in tbl.iter(qn("w:p")):
                text = DB.ptext(para)
                ids = set(DB.identifier_tokens(text))
                for w in text.split():
                    words.update([w] if w in ids or "-" not in w else w.split("-"))
        words = {w for w in words if len(w) >= 3}
        pdf_dir = tmp_path / ("pdf-" + out.parent.parent.name)
        # a private LibreOffice profile, so a soffice running elsewhere cannot swallow the conversion
        profile = (tmp_path / "lo-profile").as_uri()
        r = subprocess.run(["soffice", f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
                            "--outdir", str(pdf_dir), str(out)],
                           capture_output=True, text=True, timeout=300)
        pdf = pdf_dir / (out.stem + ".pdf")
        assert pdf.exists(), r.stderr
        text = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True,
                              text=True).stdout.replace("‑", "-")
        assert sorted(w for w in words if w not in text) == []


# ------------------------------------------------- house options and captions

def test_date_format_from_the_house(tmp_path):
    prof = northwind_profile()
    prof.setdefault("house", {})["date_format"] = "iso"
    root = factory.write_project(tmp_path / "doc", prof)
    st = DB.read_structure(DB.build(Project.load(root))[0])
    assert st.tables[1][1][0] == "2026-10-01"
    assert DB.format_date("2026-10-17") == "17 October 2026"
    assert DB.format_date("2026-10-17", "%d/%m/%Y") == "17/10/2026"
    assert DB.format_date("TBC") == "TBC"


def test_view_caption_names_the_metamodel_and_takes_a_note(tmp_path):
    prof = northwind_profile()
    prof["systems"]["ea_viewpoint_names"] = {"solution-context": "Context Model"}
    prof.setdefault("house", {})["captions"] = {"view_note": "Drawn from the model in {sec:appendix/object-guide}"}
    root = factory.write_project(tmp_path / "doc", prof)
    p = Project.load(root)
    st = DB.read_structure(DB.build(p)[0])
    n = p.number("appendix", "object-guide")
    assert st.captions[-1] == f"View 1. Parcel context (ArchiMate Context Model). Drawn from the model in {n}."
    cfg = yaml.safe_load((root / "arch-doc.yaml").read_text())
    cfg["captions"] = {"view_note": "Not yet in the repository"}
    (root / "arch-doc.yaml").write_text(yaml.safe_dump(cfg))
    st = DB.read_structure(DB.build(Project.load(root))[0])
    assert st.captions[-1].endswith("(ArchiMate Context Model). Not yet in the repository.")


def test_catalogue_shared_column(tmp_path):
    import json
    root = factory.write_project(tmp_path / "doc", northwind_profile())
    model = json.loads((root / "diagrams" / "model.json").read_text())
    model["objects"].append({"id": "A3", "type": "Application", "name": "Identity", "inheritedFrom": "Platform"})
    (root / "diagrams" / "model.json").write_text(json.dumps(model))
    cfg = yaml.safe_load((root / "arch-doc.yaml").read_text())
    cfg["catalogue"] = {"columns": ["shared"]}
    (root / "arch-doc.yaml").write_text(yaml.safe_dump(cfg))
    st = DB.read_structure(DB.build(Project.load(root))[0])
    assert [["Object type", "Objects in this design", "Shared with the Platform"], ["Application", "2", "1"],
            ["Data Object", "1", "0"]] in st.tables


def test_template_stakeholder_table_maps_business_unit_to_team(tmp_path):
    from docx import Document as D
    path = tmp_path / "t.docx"
    factory.make_template(path)
    d = D(str(path))
    t = d.tables[1]
    t._tbl.addnext(factory._table(d, ["Role", "Name", "Business Unit", "Interest"], ["", "", "", ""])._tbl)
    d.save(str(path))
    p, out = fill(tmp_path, path, FIELDS)
    st = DB.read_structure(out)
    assert ["Sponsor", "Head of Parcels", "Parcels", "Delivery on time"] in \
        next(t for t in st.tables if t[0] == ["Role", "Name", "Business Unit", "Interest"])


def test_template_mode_never_adds_sections(tmp_path, template):
    content = {"19-risks.md": "# risks\n## architecture-risks\n"
               "REGISTER risk cols id,title,category,cause,consequence,likelihood,impact,rating,mitigation,"
               "residual,owner,status\n"}
    p, out = fill(tmp_path, template, FIELDS, content=content)
    st = DB.read_structure(out)
    assert st.sections == DB.read_structure(template).sections


def test_register_rows_name_an_inherited_owner_once(tmp_path):
    content = {"19-risks.md": "# risks\n## architecture-risks\nREGISTER risk cols id,title,likelihood,impact,owner\n"}
    root = factory.write_project(tmp_path / "doc", northwind_profile(), content=content)
    (root / "registers" / "risk.yaml").write_text(
        "- {id: RSK-01, title: Limits, category: architecture, likelihood: TBC, impact: TBC, "
        "owner: Integration lead, status: open}\n")
    p = Project.load(root)
    block = next(b for b in p.blocks() if getattr(b, "type", None) == "risk")
    out = DB.build(p)[0]
    row = next(r for t in DB.read_structure(out).tables for r in t if r[:1] == ["RSK-01"])
    assert row == block.compact_rows[0]
    assert sum("Integration lead" in c for c in row) == 1  # the owner, named once


# ------------------------------------------------- template guidance marks

def _mark_template(path):
    """Strike "Physical " in front of a chapter heading and highlight it, and
    strike and highlight the sample cells of a template table."""
    from docx import Document as D
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    d = D(str(path))
    for p in d.paragraphs:
        if p.text == "Solution Context & Blueprint Alignment":
            p._p.insert(1 if p._p.pPr is not None else 0, parse_xml(
                f'<w:r {nsdecls("w")}><w:rPr><w:strike/><w:highlight w:val="yellow"/></w:rPr>'
                f'<w:t xml:space="preserve">Physical </w:t></w:r>'))
            for r in p.runs[1:]:
                r.font.highlight_color = 7  # yellow
    for t in d.tables:
        if [c.text for c in t.rows[0].cells] == ["Scope Item", "Description"]:
            for c in t.rows[0].cells + t.rows[1].cells:
                for r in c.paragraphs[0].runs:
                    r.font.strike = True
                    r.font.highlight_color = 7
                    r.font.italic = True
    d.save(str(path))
    return path


def test_template_guidance_marks_never_reach_new_text(tmp_path, template):
    _mark_template(template)
    p, out = fill(tmp_path, template, FIELDS)
    d = Document(str(out))
    body = d.element.body
    assert body.find(".//" + qn("w:highlight")) is None
    struck = [x for x in body.iter(qn("w:strike")) if x.get(qn("w:val")) not in ("0", "false")]
    assert struck == []
    heading = next(para for para in d.paragraphs if para.text == p.chapter("context").title)
    assert all(r._r.rPr is None or r._r.rPr.find(qn("w:strike")) is None for r in heading.runs)
    scope = next(t for t in d.tables if [c.text for c in t.rows[0].cells] == ["Scope Item", "Description"])
    for c in scope.rows[1].cells:
        for r in c.paragraphs[0].runs:
            assert not r.italic and not r.font.strike and r.font.highlight_color is None


def test_header_rows_repeat_and_keep_with_the_first_row(neutral, tmp_path, template):
    p, out = neutral
    _, filled = fill(tmp_path, template, FIELDS)
    for path in (out, filled):
        d = Document(str(path))
        for t in d.tables:
            if len(t.rows) < 2:
                continue
            tr = t.rows[0]._tr
            assert tr.trPr is not None and tr.trPr.find(qn("w:tblHeader")) is not None
            assert all(par.paragraph_format.keep_with_next for c in t.rows[0].cells for par in c.paragraphs)
    # the house header's look is written into the row, so it survives the keep-with-next
    d = Document(str(out))
    risk = next(t for t in d.tables if t.rows[0].cells[0].text == "ID" and t.rows[0].cells[1].text == "Risk")
    run = risk.rows[0].cells[1].paragraphs[0].runs[0]
    assert run.bold and str(run.font.color.rgb) == "FFFFFF"


def test_a_table_too_wide_for_the_template_is_named(tmp_path, template, capsys):
    cols = ",".join(["id", "title", "category", "cause", "consequence", "likelihood", "impact", "rating",
                     "mitigation", "residual", "owner", "status"])
    content = {"19-risks.md": f"# risks\n## architecture-risks\nREGISTER risk cols {cols}\n"}
    root = factory.write_project(tmp_path / "wide", northwind_profile(), template=str(template), content=content)
    (root / "registers" / "risk.yaml").write_text(
        "- {id: RSK-01, title: Throttling, category: architecture, cause: Carrier-imposed-ratelimits-everywhere, "
        "consequence: Delayed-notifications-to-customers, likelihood: medium, impact: high, rating: high, "
        "mitigation: Queueing-and-backoff, residual: medium, owner: Integration-lead, status: mitigating}\n")
    p = Project.load(root)
    DB.build(p)
    err = capsys.readouterr().err
    n = p.number("risks", "architecture-risks")
    assert f"the risk register s{n.replace('.', '-')}-t1 in {n} Architecture Risks" in err
    assert "template's page width" in err and "cols" in err


def test_catalogue_headers_take_labels(tmp_path):
    root = factory.write_project(tmp_path / "doc", northwind_profile())
    cfg = yaml.safe_load((root / "arch-doc.yaml").read_text())
    cfg["catalogue"] = {"labels": {"type": "Model object type", "objects": "{doc:short} objects", "examples": "Such as"}}
    (root / "arch-doc.yaml").write_text(yaml.safe_dump(cfg))
    st = DB.read_structure(DB.build(Project.load(root))[0])
    assert ["Model object type", "PT objects", "Such as"] in [t[0] for t in st.tables]
